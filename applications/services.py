from docx import Document
from docx.shared import Pt
from datetime import datetime
import os
from django.conf import settings

class WordDocumentGenerator:
    def __init__(self, application):
        self.application = application
    
    def fill_application_template(self):
        """Заполняет готовый шаблон Word данными из заявления"""
        
        # Путь к шаблону
        template_path = os.path.join(settings.BASE_DIR, 'templates', 'applications', 'Заявление.docx')
        
        # Проверяем существование шаблона
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Шаблон не найден: {template_path}")
        
        # Открываем шаблон
        doc = Document(template_path)
        
        # Устанавливаем шрифт Times New Roman 12pt для всего документа
        self._set_document_font(doc)
        
        # Словарь замен
        replacements = {
            '{{ФИО_родителя}}': self._get_parent_name(),
            '{{ФИО_ребенка}}': self.application.child_full_name,
            '{{дата_рождения_ребенка}}': self.application.child_birth_date.strftime("%d.%m.%Y"),
            '{{адрес_регистрации}}': self.application.registration_address,
            '{{фактический_адрес_проживания}}': self.application.actual_address,
            '{{серия_свидетельства}}': self.application.birth_certificate_series,
            '{{номер_свидетельства}}': self.application.birth_certificate_number,
            '{{выдано}}': self.application.birth_certificate_issued_by or '',
            '{{ФИО_матери}}': self.application.mother_full_name,
            '{{номер_телефона}}': self.application.mother_phone,
            '{{ФИО_отца}}': self.application.father_full_name,
            '{{дата}}': datetime.now().strftime("%d.%m.%Y")
        }
        
        # Заменяем текст во всем документе
        self._replace_text_in_document(doc, replacements)
        
        # Сохраняем заполненный документ
        output_filename = f"Заявление_{self.application.child_full_name.replace(' ', '_')}.docx"
        output_path = os.path.join(settings.MEDIA_ROOT, 'applications', 'generated', output_filename)
        
        # Создаем директорию если не существует
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        # Сохраняем документ
        doc.save(output_path)
        
        return os.path.join('applications', 'generated', output_filename)
    
    def _set_document_font(self, doc):
        """Устанавливает шрифт Times New Roman 12pt для всего документа"""
        # Устанавливаем для стиля Normal
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        
        # Также устанавливаем для других основных стилей
        try:
            heading_style = doc.styles['Heading 1']
            heading_font = heading_style.font
            heading_font.name = 'Times New Roman'
            heading_font.size = Pt(14)  # Заголовки чуть больше
        except:
            pass
        
        try:
            heading2_style = doc.styles['Heading 2']
            heading2_font = heading2_style.font
            heading2_font.name = 'Times New Roman'
            heading2_font.size = Pt(12)
        except:
            pass
    
    def _replace_text_in_document(self, doc, replacements):
        """Заменяет текст во всем документе и устанавливает шрифт"""
        # Заменяем в параграфах
        for paragraph in doc.paragraphs:
            self._replace_in_paragraph_with_font(paragraph, replacements)
        
        # Заменяем в таблицах
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph_with_font(paragraph, replacements)
    
    def _replace_in_paragraph_with_font(self, paragraph, replacements):
        """Заменяет текст в параграфе и устанавливает шрифт Times New Roman 12pt"""
        original_text = paragraph.text
        
        # Проверяем, есть ли плейсхолдеры для замены
        has_placeholders = False
        for key in replacements.keys():
            if key in original_text:
                has_placeholders = True
                break
        
        if not has_placeholders:
            # Все равно устанавливаем шрифт для всего текста
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
            return
        
        # Создаем новый текст с замененными значениями
        new_text = original_text
        for key, value in replacements.items():
            new_text = new_text.replace(key, value)
        
        # Если текст изменился, обновляем параграф
        if new_text != original_text:
            # Полностью очищаем параграф
            paragraph.clear()
            # Добавляем новый текст с нужным шрифтом
            new_run = paragraph.add_run(new_text)
            new_run.font.name = 'Times New Roman'
            new_run.font.size = Pt(12)
        else:
            # Если замен не было, просто устанавливаем шрифт для существующих runs
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
    
    def _get_parent_name(self):
        """Определяет ФИО родителя для подписи"""
        if self.application.mother_full_name:
            return self.application.mother_full_name
        elif self.application.father_full_name:
            return self.application.father_full_name
        else:
            return self.application.parent.user.get_full_name()
