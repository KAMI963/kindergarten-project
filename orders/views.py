# orders/views.py

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.db import transaction, models
from datetime import date, datetime
from docxtpl import DocxTemplate
from django.conf import settings
import json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import io
import os
import base64
import traceback
from django.views.decorators.csrf import csrf_exempt

from .models import EnrollmentOrder
from .forms import GroupedBulkEnrollmentForm, EnrollmentOrderForm
from children.models import Child, Group, ChildParent
from applications.models import ApplicationStatus, ChildApplication
from accounts.models import CustomUser
from contracts.models import EducationContract


# Декоратор для проверки прав заведующей
def director_required(view_func):
    decorated_view_func = user_passes_test(
        lambda u: u.is_authenticated and u.role == 'director',
        login_url='/accounts/login/'
    )(view_func)
    return decorated_view_func


@login_required
@director_required
def orders_list(request):
    """Список всех приказов с фильтрацией"""
    try:
        orders = EnrollmentOrder.objects.all().select_related(
            'child', 'group', 'created_by'
        ).order_by('-order_date', '-created_at')
        
        # Фильтрация
        child_name = request.GET.get('child_name')
        order_number = request.GET.get('order_number')
        order_date = request.GET.get('order_date')
        group_id = request.GET.get('group')
        
        if child_name:
            orders = orders.filter(child__full_name__icontains=child_name)
        if order_number:
            orders = orders.filter(order_number__icontains=order_number)
        if order_date:
            orders = orders.filter(order_date=order_date)
        if group_id:
            orders = orders.filter(group_id=group_id)
        
        # Статистика
        total_orders = orders.count()
        active_orders = orders.filter(is_active=True).count()
        
        # Упрощенная статистика по группам
        groups_with_stats = []
        groups = Group.objects.all()
        
        for group in groups:
            children_count = group.child_set.count()
            free_spaces = max(0, group.capacity - children_count)
            
            groups_with_stats.append({
                'group': group,
                'children_count': children_count,
                'free_spaces': free_spaces,
            })
        
        # Заявления за сегодня
        today_applications = ChildApplication.objects.filter(
            created_at__date=date.today()
        ).count()
        
        return render(request, 'orders/orders_list.html', {
            'orders': orders,
            'total_orders': total_orders,
            'active_orders': active_orders,
            'groups_with_stats': groups_with_stats,
            'today_applications': today_applications,
            'groups': groups,
            'filter_params': {
                'child_name': child_name or '',
                'order_number': order_number or '',
                'order_date': order_date or '',
                'group_id': group_id or ''
            }
        })
        
    except Exception as e:
        messages.error(request, f'Ошибка при загрузке списка приказов: {str(e)}')
        return render(request, 'orders/orders_list.html', {
            'orders': [],
            'total_orders': 0,
            'active_orders': 0,
            'groups_with_stats': [],
            'today_applications': 0,
            'groups': Group.objects.all(),
            'filter_params': {}
        })


@login_required
@director_required
def order_detail(request, pk):
    """Детальная страница приказа"""
    try:
        order = get_object_or_404(EnrollmentOrder.objects.select_related(
            'child', 'group', 'created_by'
        ), pk=pk)
        
        return render(request, 'orders/order_detail.html', {
            'order': order
        })
        
    except Exception as e:
        messages.error(request, f'Ошибка при загрузке приказа: {str(e)}')
        return redirect('orders:orders_list')


@login_required
@director_required
def order_edit(request, pk):
    """Редактирование приказа"""
    try:
        order = get_object_or_404(EnrollmentOrder, pk=pk)
        
        if request.method == 'POST':
            form = EnrollmentOrderForm(request.POST, instance=order, request=request)
            if form.is_valid():
                try:
                    form.save()
                    messages.success(request, f'Приказ №{order.order_number} успешно обновлен!')
                    return redirect('orders:order_detail', pk=order.pk)
                except Exception as e:
                    messages.error(request, f'Ошибка при обновлении приказа: {str(e)}')
            else:
                messages.error(request, 'Пожалуйста, исправьте ошибки в форме.')
        else:
            form = EnrollmentOrderForm(instance=order, request=request)
        
        return render(request, 'orders/order_form.html', {
            'form': form,
            'order': order,
            'title': f'Редактирование приказа №{order.order_number}',
            'submit_text': 'Сохранить изменения'
        })
        
    except Exception as e:
        messages.error(request, f'Ошибка при загрузке формы редактирования: {str(e)}')
        return redirect('orders:orders_list')


@login_required
@director_required
def order_delete(request, pk):
    """Удаление приказа"""
    try:
        order = get_object_or_404(EnrollmentOrder, pk=pk)
        
        if request.method == 'POST':
            try:
                order_number = order.order_number
                order.delete()
                messages.success(request, f'Приказ №{order_number} успешно удален!')
                return redirect('orders:orders_list')
            except Exception as e:
                messages.error(request, f'Ошибка при удалении приказа: {str(e)}')
                return redirect('orders:order_detail', pk=pk)
        
        return render(request, 'orders/order_confirm_delete.html', {
            'order': order
        })
        
    except Exception as e:
        messages.error(request, f'Ошибка при удалении приказа: {str(e)}')
        return redirect('orders:orders_list')


@login_required
@director_required
def order_create(request):
    """Создание нового приказа (для существующих детей)"""
    try:
        if request.method == 'POST':
            form = EnrollmentOrderForm(request.POST, request=request)
            if form.is_valid():
                try:
                    order = form.save(commit=False)
                    order.created_by = request.user
                    order.save()
                    messages.success(request, f'Приказ №{order.order_number} успешно создан!')
                    return redirect('orders:order_detail', pk=order.pk)
                except Exception as e:
                    messages.error(request, f'Ошибка при создании приказа: {str(e)}')
            else:
                messages.error(request, 'Пожалуйста, исправьте ошибки в форме.')
        else:
            form = EnrollmentOrderForm(request=request)
        
        return render(request, 'orders/order_form.html', {
            'form': form,
            'title': 'Создание приказа о зачислении',
            'submit_text': 'Создать приказ'
        })
        
    except Exception as e:
        messages.error(request, f'Ошибка при создании приказа: {str(e)}')
        return redirect('orders:orders_list')


# ==================== ФУНКЦИЯ СКЛОНЕНИЯ ФИО ====================

def decline_fio_accusative(full_name):
    """
    Склоняет ФИО для использования в документах (винительный падеж)
    Пример: Ковалева Софья Дмитриевна -> Ковалеву Софью Дмитриевну
    (зачислить КОГО?)
    """
    if not full_name:
        return full_name
    
    parts = full_name.strip().split()
    if len(parts) >= 3:
        last_name, first_name, middle_name = parts[0], parts[1], parts[2]
    elif len(parts) == 2:
        last_name, first_name = parts[0], parts[1]
        middle_name = ''
    else:
        return full_name
    
    # Склонение фамилии (винительный падеж)
    declined_last = last_name
    if last_name.endswith('а'):
        declined_last = last_name[:-1] + 'у'
    elif last_name.endswith('я'):
        declined_last = last_name[:-1] + 'ю'
    elif last_name.endswith('ий'):
        declined_last = last_name[:-2] + 'его'
    elif last_name.endswith('ой') or last_name.endswith('ый'):
        declined_last = last_name[:-2] + 'ого'
    else:
        declined_last = last_name + 'а'
    
    # Склонение имени (винительный падеж)
    declined_first = first_name
    if first_name.endswith('а'):
        declined_first = first_name[:-1] + 'у'
    elif first_name.endswith('я'):
        declined_first = first_name[:-1] + 'ю'
    elif first_name.endswith('й'):
        declined_first = first_name[:-1] + 'я'
    else:
        declined_first = first_name + 'а'
    
    # Склонение отчества (винительный падеж)
    declined_middle = middle_name
    if middle_name:
        if middle_name.endswith('а'):
            declined_middle = middle_name[:-1] + 'у'
        elif middle_name.endswith('я'):
            declined_middle = middle_name[:-1] + 'ю'
        elif middle_name.endswith('на'):
            declined_middle = middle_name[:-2] + 'ну'
        elif middle_name.endswith('чна'):
            declined_middle = middle_name[:-3] + 'чну'
        else:
            declined_middle = middle_name + 'а'
    
    if middle_name:
        return f"{declined_last} {declined_first} {declined_middle}"
    else:
        return f"{declined_last} {declined_first}"

@login_required
@director_required
def generate_single_order_word(request, pk):
    """Генерация приказа в формате Word для одного ребенка"""
    try:
        order = get_object_or_404(EnrollmentOrder.objects.select_related(
            'child', 'group', 'created_by', 'application'
        ), pk=pk)
        
        # Путь к шаблону
        template_path = os.path.join(settings.BASE_DIR, 'templates', 'documents', 'Приказ_о_зачислении_МБДОУ_РЯБИНУШКА.docx')
        
        # Проверяем существование файла
        if not os.path.exists(template_path):
            messages.error(request, f'Шаблон не найден по пути: {template_path}')
            return redirect('orders:order_detail', pk=pk)
        
        # Получаем ФИО ребенка
        child_full_name = order.child.full_name
        print(f"Оригинальное ФИО: {child_full_name}")
        
        # Склоняем ФИО ребенка в ВИНИТЕЛЬНЫЙ падеж
        child_name_declined = decline_fio_accusative(child_full_name)
        print(f"Склоненное ФИО (винительный падеж): {child_name_declined}")
        
        # Формируем даты
        order_date_str = order.order_date.strftime('%d.%m.%Y')
        enrollment_date_str = order.enrollment_date.strftime('%d.%m.%Y')
        
        # Получаем фамилию для имени файла
        last_name = child_full_name.split()[0] if child_full_name.split() else "ребенок"
        
        # Данные для свидетельства
        birth_certificate = ''
        issued_by = ''
        issue_date = ''
        
        if order.application:
            series = order.application.birth_certificate_series or ''
            number = order.application.birth_certificate_number or ''
            birth_certificate = f"{series} {number}".strip() if (series or number) else 'не указано'
            issued_by = order.application.birth_certificate_issued_by or 'не указан'
            if order.application.birth_certificate_issue_date:
                issue_date = order.application.birth_certificate_issue_date.strftime('%d.%m.%Y')
        
        # Данные для шаблона
        context = {
            'дата_приказа': order_date_str,
            'номер_приказа': order.order_number,
            'дата_зачисления': enrollment_date_str,
            'название_группы': order.group.name,
            'дети': [
                {
                    'номер': 1,
                    'ФИО_ребенка': child_name_declined,
                    'дата_рождения_ребенка': order.child.birth_date.strftime('%d.%m.%Y') if order.child.birth_date else '',
                    'данные_свидетельства': birth_certificate,
                    'орган_выдачи': issued_by,
                    'дата_выдачи': issue_date
                }
            ]
        }
        
        print("Контекст для шаблона:", context)
        
        # Загружаем и заполняем шаблон
        doc = DocxTemplate(template_path)
        doc.render(context)
        
        # Сохраняем в буфер
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        # Формируем имя файла на русском с правильной кодировкой
        filename = f"Приказ_о_зачислении_{order.order_number}_{last_name}.docx"
        
        # Импортируем quote для кодирования
        from urllib.parse import quote
        
        # Создаем ответ с правильной кодировкой для русских букв
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        # Кодируем имя файла для поддержки кириллицы
        encoded_filename = quote(filename)
        response['Content-Disposition'] = f"attachment; filename*=UTF-8''{encoded_filename}"
        
        return response
        
    except Exception as e:
        messages.error(request, f'Ошибка при генерации документа: {str(e)}')
        import traceback
        traceback.print_exc()
        return redirect('orders:order_detail', pk=pk)
        

def create_document_from_scratch(context):
    """Создание документа Word с нуля, если нет шаблона"""
    document = Document()
    
    # Настройка стилей
    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    
    # Заголовок учреждения
    title = document.add_heading('Муниципальное бюджетное дошкольное образовательное учреждение', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.size = Pt(14)
    
    subtitle = document.add_heading('«Карабашский детский сад общеразвивающего вида №1 «Рябинушка»»', level=2)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(14)
    
    address = document.add_paragraph('423229, Республика Татарстан, Бугульминский район, пгт. Карабаш, ул. Октябрьская, д. 6')
    address.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if address.runs:
        address.runs[0].font.size = Pt(12)
    
    document.add_paragraph()
    
    # ПРИКАЗ
    order_title = document.add_heading('ПРИКАЗ', level=1)
    order_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in order_title.runs:
        run.font.size = Pt(16)
        run.bold = True
    
    order_subtitle = document.add_paragraph('О зачислении в МБДОУ «Рябинушка»')
    order_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if order_subtitle.runs:
        order_subtitle.runs[0].font.size = Pt(14)
    
    document.add_paragraph()
    
    # Дата и номер приказа
    order_info = document.add_paragraph()
    order_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    order_info.add_run(f'от «{context.get("order_date", "")}»\n').font.size = Pt(14)
    order_info.add_run(f'№ {context.get("order_number", "")}').font.size = Pt(14)
    
    document.add_paragraph()
    
    # Основание
    basis_para = document.add_paragraph()
    basis_para.add_run('На основании:\n').bold = True
    basis_para.add_run(context.get('basis_documents', 'заявления родителей, свидетельства о рождении, медицинской карты'))
    
    document.add_paragraph()
    
    # ПРИКАЗЫВАЮ
    command_para = document.add_paragraph()
    command_para.add_run('ПРИКАЗЫВАЮ:').bold = True
    if command_para.runs:
        command_para.runs[0].font.size = Pt(14)
    
    document.add_paragraph()
    
    # Текст приказа
    order_text = document.add_paragraph()
    order_text.add_run(f'1. Зачислить в МБДОУ «Рябинушка» с «{context.get("enrollment_date", "")}» '
                       f'в группу {context.get("group_name", "")} ребенка:\n')
    
    if context.get('child_full_name') and context.get('child_birth_date'):
        order_text.add_run(f'   {context["child_full_name"]}, {context["child_birth_date"]} г.р.\n\n')
    else:
        # Для массового зачисления
        if context.get('дети'):
            for child in context['дети']:
                order_text.add_run(f'   {child["ФИО_ребенка"]}, {child["дата_рождения_ребенка"]} г.р.\n')
            order_text.add_run('\n')
    
    order_text.add_run('2. Заведующей обеспечить:\n'
                       '   - организацию учебно-воспитательного процесса;\n'
                       '   - ведение личного дела воспитанника.\n\n')
    order_text.add_run('3. Контроль за исполнением приказа оставляю за собой.')
    
    document.add_paragraph()
    document.add_paragraph()
    
    # Подпись
    signature_table = document.add_table(rows=2, cols=2)
    signature_table.autofit = False
    signature_table.columns[0].width = Inches(3)
    signature_table.columns[1].width = Inches(3)
    
    # Левая колонка
    left_cell = signature_table.cell(0, 0)
    left_cell.text = 'Заведующая МБДОУ «Рябинушка»'
    left_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Правая колонка
    right_cell = signature_table.cell(0, 1)
    right_cell.text = '_________________________ Н.В. Михайлова'
    right_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Вторая строка - дата
    date_cell = signature_table.cell(1, 0)
    date_cell.text = f'«{context.get("order_date", "")}»'
    date_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    document.add_paragraph()
    document.add_paragraph()
    
    # Ознакомление
    familiar = document.add_paragraph('С приказом ознакомлен(а):\n')
    familiar.add_run('_________________________')
    
    return document


# ==================== ГРУППОВОЕ ЗАЧИСЛЕНИЕ ====================

@login_required
@director_required
def grouped_bulk_enrollment(request):
    """Массовое зачисление по группам"""
    from django.db import connection
    from applications.models import ApplicationStatus
    
    # GET запрос - показываем форму
    if request.method == 'GET':
        groups = Group.objects.all()
        
        # ИСПРАВЛЕНО: используем правильные статусы
        allowed_statuses = [ApplicationStatus.APPROVED, ApplicationStatus.QUEUE, 
                           ApplicationStatus.INVITED, ApplicationStatus.PROCESSING]
        
        total_queue = ChildApplication.objects.filter(
            status__in=allowed_statuses
        ).count()
        
        print(f"Найдено заявлений со статусами {allowed_statuses}: {total_queue}")
        
        # Выводим все заявления для отладки
        all_apps = ChildApplication.objects.all()
        print("Все заявления в БД:")
        for app in all_apps:
            print(f"  ID: {app.id}, Имя: {app.child_full_name}, Статус: {app.status}")
        
        groups_stats = []
        for group in groups:
            children_count = group.child_set.count()
            free_spaces = max(0, group.capacity - children_count)
            groups_stats.append({
                'group': group,
                'children_count': children_count,
                'free_spaces': free_spaces,
                'capacity': group.capacity,
            })
        
        return render(request, 'orders/grouped_bulk_enrollment.html', {
            'groups': groups,
            'groups_stats': groups_stats,
            'total_approved': total_queue,
            'today': date.today(),
        })
    
    # POST запрос - обрабатываем создание приказов
    if request.method == 'POST':
        try:
            group_id = request.POST.get('group_id')
            order_date_str = request.POST.get('order_date')
            enrollment_date_str = request.POST.get('enrollment_date')
            basis_documents = request.POST.get('basis_documents')
            application_ids = request.POST.getlist('applications')
            
            if not group_id:
                messages.error(request, 'Выберите группу')
                return redirect('orders:grouped_bulk_enrollment')
            
            if not application_ids:
                messages.error(request, 'Выберите хотя бы одного ребенка')
                return redirect('orders:grouped_bulk_enrollment')
            
            group = get_object_or_404(Group, id=group_id)
            order_date = datetime.strptime(order_date_str, '%Y-%m-%d').date() if order_date_str else date.today()
            enrollment_date = datetime.strptime(enrollment_date_str, '%Y-%m-%d').date() if enrollment_date_str else date.today()
            
            allowed_statuses = ['invited', 'approved', 'queue', 'processing']
            applications = ChildApplication.objects.filter(
                id__in=application_ids,
                status__in=allowed_statuses
            )
            
            if not applications.exists():
                messages.error(request, 'Выбранные заявления не найдены или уже зачислены')
                return redirect('orders:grouped_bulk_enrollment')
            
            free_spaces = group.capacity - group.child_set.count()
            if len(applications) > free_spaces:
                messages.error(request, f'В группе только {free_spaces} свободных мест')
                return redirect('orders:grouped_bulk_enrollment')
            
            created_orders = []
            errors = []
            
            for app in applications:
                child_id = None
                try:
                    print(f"\n=== Обработка: {app.child_full_name} ===")
                    
                    with connection.cursor() as cursor:
                        cursor.execute("""
                            INSERT INTO children_child (
                                full_name, birth_date, gender, blood_type,
                                allergies, chronic_diseases, special_needs,
                                registration_address, actual_address,
                                is_active, created_at, updated_at,
                                birth_certificate, birth_certificate_issued_by,
                                disability_certificate, medical_card_number,
                                policy_oms_number, snils,
                                consent_data_processing, consent_medical_intervention,
                                consent_photo_video, group_id, enrollment_date
                            ) VALUES (
                                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
                            )
                            RETURNING id
                        """, [
                            app.child_full_name,
                            app.child_birth_date,
                            getattr(app, 'child_gender', 'M'),
                            'I',
                            getattr(app, 'allergies', '') or 'Нет',
                            getattr(app, 'chronic_diseases', '') or 'Нет',
                            'Нет',
                            getattr(app, 'registration_address', ''),
                            getattr(app, 'actual_address', ''),
                            True,
                            timezone.now(),
                            timezone.now(),
                            f"{getattr(app, 'birth_certificate_series', '') or ''} {getattr(app, 'birth_certificate_number', '') or ''}".strip() or f"Временное-{app.id}",
                            getattr(app, 'birth_certificate_issued_by', '') or 'Не указано',
                            '',
                            '',
                            '',
                            getattr(app, 'child_snils', '') or '',
                            True,
                            True,
                            True,
                            group.id,
                            enrollment_date
                        ])
                        child_id = cursor.fetchone()[0]
                        print(f"  ✓ Ребенок создан: ID={child_id}")
                    
                    if hasattr(app, 'parent') and app.parent:
                        with connection.cursor() as cursor:
                            cursor.execute("""
                                INSERT INTO children_childparent (
                                    child_id, parent_id, relation, is_primary
                                ) VALUES (%s, %s, %s, %s)
                                ON CONFLICT (child_id, parent_id) DO NOTHING
                            """, [child_id, app.parent.id, 'mother', True])
                            print(f"  ✓ Связь с родителем создана")
                    
                    # Генерируем номер приказа
                    year = order_date.year
                    with connection.cursor() as cursor:
                        cursor.execute("""
                            SELECT order_number FROM orders_enrollmentorder 
                            WHERE order_date >= %s AND order_date <= %s
                            ORDER BY id DESC LIMIT 1
                        """, [f"{year}-01-01", f"{year}-12-31"])
                        last_order = cursor.fetchone()
                        
                        if last_order and last_order[0]:
                            try:
                                last_num = int(last_order[0].split('-')[-1])
                                new_num = last_num + 1
                            except:
                                new_num = 1
                        else:
                            new_num = 1
                    
                    order_number = f"{year}-{new_num:04d}"
                    print(f"  ✓ Номер приказа: {order_number}")
                    
                    # Вставляем приказ
                    with connection.cursor() as cursor:
                        cursor.execute("""
                            INSERT INTO orders_enrollmentorder (
                                order_number, order_date, enrollment_date, child_id,
                                group_id, created_by_id, application_id, basis_documents,
                                status, is_active, created_at, updated_at, order_type
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                        """, [
                            order_number, order_date, enrollment_date, child_id,
                            group.id, request.user.id, app.id, basis_documents,
                            'issued', True, timezone.now(), timezone.now(), 'enrollment'
                        ])
                        print(f"  ✓ Приказ создан: №={order_number}")
                        created_orders.append({'number': order_number, 'child_name': app.child_full_name})
                    
                    # Обновляем статус заявления
                    with connection.cursor() as cursor:
                        cursor.execute("""
                            UPDATE applications_childapplication 
                            SET status = %s, 
                                enrolled_group = %s,
                                enrollment_order_number = %s,
                                enrollment_order_date = %s,
                                updated_at = %s
                            WHERE id = %s
                        """, ['enrolled', group.name, order_number, order_date, timezone.now(), app.id])
                    print(f"  ✓ Статус заявления обновлен")
                    
                except Exception as e:
                    error_msg = f"{app.child_full_name}: {str(e)}"
                    errors.append(error_msg)
                    print(f"  ✗ ОШИБКА: {error_msg}")
                    traceback.print_exc()
                    
                    if child_id:
                        try:
                            with connection.cursor() as cursor:
                                cursor.execute("DELETE FROM children_child WHERE id = %s", [child_id])
                            print(f"  ✓ Откат: ребенок {child_id} удален")
                        except:
                            pass
                    continue
            
            if created_orders:
                messages.success(request, f'✅ Успешно создано {len(created_orders)} приказов!')
            
            if errors:
                messages.warning(request, f'⚠️ Ошибки при создании ({len(errors)}):<br>' + '<br>'.join(errors[:3]))
            
            return redirect('orders:orders_list')
            
        except Exception as e:
            print(f"Критическая ошибка: {str(e)}")
            traceback.print_exc()
            messages.error(request, f'Произошла ошибка: {str(e)}')
            return redirect('orders:grouped_bulk_enrollment')
    
    return redirect('orders:orders_list')


def generate_order_number(order_date):
    """Генерация номера приказа"""
    year = order_date.year
    last_order = EnrollmentOrder.objects.filter(
        order_date__year=year
    ).order_by('-id').first()
    
    if last_order and last_order.order_number:
        try:
            parts = last_order.order_number.split('-')
            if len(parts) == 2:
                last_num = int(parts[1])
                new_num = last_num + 1
            else:
                new_num = 1
        except (ValueError, IndexError):
            new_num = 1
    else:
        new_num = 1
    
    return f"{year}-{new_num:04d}"


def get_suitable_applications_for_group(group):
    """Получает подходящие заявления для группы"""
    from datetime import date
    
    print(f"=== get_suitable_applications_for_group для группы {group.name} ===")
    
    applications_with_orders = EnrollmentOrder.objects.filter(
        application__isnull=False
    ).values_list('application_id', flat=True)
    
    # ИСПРАВЛЕНО: используем правильные значения статусов из модели
    allowed_statuses = [
        ApplicationStatus.APPROVED,  # 'approved'
        ApplicationStatus.QUEUE,      # 'queue'
        ApplicationStatus.INVITED,    # 'invited'
        ApplicationStatus.PROCESSING, # 'processing'
    ]
    
    # Также можно добавить строковые значения для безопасности
    allowed_status_strings = ['approved', 'queue', 'invited', 'processing']
    
    eligible_applications = ChildApplication.objects.filter(
        models.Q(status__in=allowed_statuses) | models.Q(status__in=allowed_status_strings)
    ).exclude(
        id__in=applications_with_orders
    )
    
    print(f"Найдено заявлений с допустимыми статусами: {eligible_applications.count()}")
    
    # Выводим все статусы для отладки
    all_statuses = ChildApplication.objects.values_list('status', flat=True).distinct()
    print(f"Все статусы в БД: {list(all_statuses)}")
    
    for app in eligible_applications[:5]:
        print(f"  - {app.child_full_name}: статус '{app.status}'")
    
    suitable_applications = []
    today = date.today()
    
    for app in eligible_applications:
        age = calculate_age(app.child_birth_date, today)
        
        if is_age_appropriate_for_group(age, group):
            suitable_applications.append(app)
            print(f"✓ ПОДХОДИТ: {app.child_full_name}, возраст: {age}, статус: '{app.status}'")
        else:
            print(f"✗ НЕ ПОДХОДИТ ПО ВОЗРАСТУ: {app.child_full_name}, возраст: {age}")
    
    print(f"Итого подходящих: {len(suitable_applications)}")
    return suitable_applications


@login_required
@director_required
def get_applications_for_group(request, group_id):
    """AJAX-запрос для получения заявлений для выбранной группы"""
    try:
        group = Group.objects.get(id=group_id)
        suitable_applications = get_suitable_applications_for_group(group)
        
        applications_data = []
        today = date.today()
        
        for app in suitable_applications:
            age = calculate_age(app.child_birth_date, today)
            applications_data.append({
                'id': app.id,
                'child_full_name': app.child_full_name,
                'child_birth_date': app.child_birth_date.strftime('%d.%m.%Y'),
                'age': f"{age:.1f}",
                'birth_certificate_series': app.birth_certificate_series or '',
                'birth_certificate_number': app.birth_certificate_number or '',
                'birth_certificate_issued_by': app.birth_certificate_issued_by or '',
                'birth_certificate_issue_date': app.birth_certificate_issue_date.strftime('%d.%m.%Y') if app.birth_certificate_issue_date else 'не указана'
            })
        
        children_count = group.child_set.count()
        free_spaces = max(0, group.capacity - children_count)
        
        return JsonResponse({
            'success': True,
            'applications': applications_data,
            'free_spaces': free_spaces,
            'children_count': children_count,
            'capacity': group.capacity
        })
        
    except Group.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Группа не найдена'})
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'success': False, 'error': str(e)})


# ==================== AJAX ПЕЧАТЬ ====================

@login_required
@director_required
@csrf_exempt
def print_order_ajax(request, pk):
    """AJAX-обработчик для печати приказа с анимацией (Word документ)"""
    try:
        order = get_object_or_404(EnrollmentOrder.objects.select_related(
            'child', 'group', 'created_by', 'application'
        ), pk=pk)
        
        response = generate_single_order_word(request, pk)
        
        if response.status_code == 200:
            doc_content = response.content
            doc_base64 = base64.b64encode(doc_content).decode('utf-8')
            
            return JsonResponse({
                'success': True,
                'message': 'Документ сформирован',
                'doc_base64': doc_base64,
                'filename': f'Приказ_{order.order_number}.docx'
            })
        else:
            return JsonResponse({
                'success': False,
                'error': 'Ошибка при формировании документа'
            })
            
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'error': str(e)
        })


# ==================== ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ====================

def calculate_age(birth_date, reference_date):
    """Рассчитывает точный возраст ребенка в годах"""
    age = reference_date.year - birth_date.year
    months_diff = reference_date.month - birth_date.month
    days_diff = reference_date.day - birth_date.day
    
    if months_diff < 0 or (months_diff == 0 and days_diff < 0):
        age -= 1
        months_diff += 12
    
    exact_age = age + (months_diff / 12.0)
    return round(exact_age, 1)


def is_age_appropriate_for_group(age, group):
    """Проверяет, подходит ли возраст ребенка для группы"""
    age_ranges = {
        'nursery': (1.5, 3.5),
        'junior': (3, 4.5),
        'middle': (4, 5.5),
        'senior': (5, 6.5),
        'preparatory': (6, 7.5),
    }
    
    if hasattr(group, 'age_category') and group.age_category:
        category = group.age_category.lower() if isinstance(group.age_category, str) else ''
        
        if category in age_ranges:
            min_age, max_age = age_ranges[category]
            return min_age <= age < max_age
    
    group_name_lower = group.name.lower()
    for key, (min_age, max_age) in age_ranges.items():
        if key in group_name_lower:
            return min_age <= age < max_age
    
    return True


@login_required
def get_child_info(request, child_id):
    """AJAX-запрос для получения информации о ребенке"""
    if request.user.role != 'director':
        return JsonResponse({'error': 'Доступ запрещен'}, status=403)
    
    try:
        child = Child.objects.select_related('group').get(id=child_id)
        
        data = {
            'full_name': child.full_name,
            'birth_date': child.birth_date.strftime('%d.%m.%Y'),
            'age': child.get_age(),
            'gender': child.get_gender_display(),
            'address': child.actual_address,
            'group_name': child.group.name if child.group else 'Не распределен',
            'parents': []
        }
        
        parents = child.parents_relations.select_related('parent__user').all()
        for parent in parents:
            data['parents'].append({
                'name': parent.parent.user.get_full_name(),
                'relation': parent.get_relation_display(),
                'phone': parent.parent.user.phone or 'Не указан',
                'email': parent.parent.user.email or 'Не указан'
            })
        
        return JsonResponse(data)
        
    except Child.DoesNotExist:
        return JsonResponse({'error': 'Ребенок не найден'}, status=404)


# ==================== ОТЛАДОЧНЫЕ ФУНКЦИИ ====================

@login_required
@director_required
def debug_check(request):
    """Отладочная страница для проверки заявлений"""
    from applications.models import ChildApplication
    from children.models import Group
    from datetime import date
    
    groups = Group.objects.all()
    
    result = {
        'groups': [],
        'all_applications': [],
        'statuses': {}
    }
    
    for app in ChildApplication.objects.all():
        status = str(app.status)
        if status not in result['statuses']:
            result['statuses'][status] = 0
        result['statuses'][status] += 1
        
        result['all_applications'].append({
            'id': app.id,
            'name': app.child_full_name,
            'birth_date': str(app.child_birth_date),
            'status': status,
            'has_order': EnrollmentOrder.objects.filter(application=app).exists()
        })
    
    for group in groups:
        children_count = group.child_set.count()
        free_spaces = group.capacity - children_count
        suitable = get_suitable_applications_for_group(group)
        
        result['groups'].append({
            'id': group.id,
            'name': group.name,
            'age_category': group.age_category if hasattr(group, 'age_category') else 'не указана',
            'capacity': group.capacity,
            'children_count': children_count,
            'free_spaces': free_spaces,
            'suitable_count': len(suitable),
            'suitable_applications': [
                {'id': app.id, 'name': app.child_full_name, 'age': calculate_age(app.child_birth_date, date.today())}
                for app in suitable[:10]
            ]
        })
    
    return JsonResponse(result, json_dumps_params={'ensure_ascii': False, 'indent': 2})


@login_required
@director_required
def debug_applications(request):
    """Отладочная страница для проверки заявлений"""
    from applications.models import ApplicationStatus, ChildApplication
    from children.models import Group
    from datetime import date
    
    all_applications = ChildApplication.objects.all()
    
    status_stats = {}
    for app in all_applications:
        status = app.status
        if isinstance(status, ApplicationStatus):
            status = status.value
        if status not in status_stats:
            status_stats[status] = 0
        status_stats[status] += 1
    
    applications_with_orders = EnrollmentOrder.objects.filter(
        application__isnull=False
    ).values_list('application_id', flat=True)
    
    available_applications = ChildApplication.objects.exclude(
        id__in=applications_with_orders
    )
    
    groups = Group.objects.all()
    
    group_matches = {}
    for group in groups:
        suitable = get_suitable_applications_for_group(group)
        group_matches[group.name] = {
            'count': len(suitable),
            'applications': [(app.child_full_name, calculate_age(app.child_birth_date, date.today())) for app in suitable[:5]]
        }
    
    return JsonResponse({
        'total_applications': all_applications.count(),
        'status_stats': status_stats,
        'available_applications': available_applications.count(),
        'applications_without_orders': [
            {
                'id': app.id,
                'name': app.child_full_name,
                'status': str(app.status),
                'birth_date': str(app.child_birth_date),
                'created_at': str(app.created_at)
            } for app in available_applications[:20]
        ],
        'group_matches': group_matches,
        'all_status_values': list(set(str(app.status) for app in all_applications))
    })


@login_required
@director_required
def generate_group_order_word(request):
    """Генерация приказа по шаблону Word для выбранной группы"""
    try:
        if request.method == 'POST':
            order_date_str = request.POST.get('order_date')
            enrollment_date_str = request.POST.get('enrollment_date')
            group_id = request.POST.get('group')
            application_ids = request.POST.getlist('applications')
            basis_documents = request.POST.get('basis_documents', 'На основании заявления родителей, свидетельства о рождении, медицинской карты')
            
            if not all([order_date_str, enrollment_date_str, group_id]):
                messages.error(request, 'Не все обязательные поля заполнены!')
                return redirect('orders:grouped_bulk_enrollment')
            
            if not application_ids:
                messages.error(request, 'Не выбрано ни одного заявления!')
                return redirect('orders:grouped_bulk_enrollment')
            
            try:
                order_date = datetime.strptime(order_date_str, '%Y-%m-%d').date()
                enrollment_date = datetime.strptime(enrollment_date_str, '%Y-%m-%d').date()
            except ValueError as e:
                messages.error(request, f'Ошибка в формате даты: {str(e)}')
                return redirect('orders:grouped_bulk_enrollment')
            
            group = Group.objects.get(id=group_id)
            applications = ChildApplication.objects.filter(id__in=application_ids, status='approved')
            
            if not applications:
                messages.error(request, 'Выбранные заявления не найдены или не одобрены!')
                return redirect('orders:grouped_bulk_enrollment')
            
            order_number = generate_order_number(order_date)
            return create_word_document(order_date, enrollment_date, order_number, group, applications, basis_documents)
        else:
            messages.error(request, 'Неверный метод запроса')
            return redirect('orders:grouped_bulk_enrollment')
            
    except Exception as e:
        messages.error(request, f'Ошибка при генерации документа: {str(e)}')
        return redirect('orders:grouped_bulk_enrollment')


def create_word_document(order_date, enrollment_date, order_number, group, applications, basis_documents):
    """Создание документа Word"""
    try:
        template_path = os.path.join(settings.BASE_DIR, 'templates', 'documents', 'Приказ_о_зачислении_МБДОУ_РЯБИНУШКА.docx')
        
        if os.path.exists(template_path):
            return create_document_from_template(template_path, order_date, enrollment_date, order_number, group, applications)
        else:
            return create_document_from_scratch({'order_date': order_date.strftime('%d.%m.%Y'), 'enrollment_date': enrollment_date.strftime('%d.%m.%Y'), 'order_number': order_number, 'group_name': group.name, 'basis_documents': basis_documents})
            
    except Exception as e:
        return create_document_from_scratch({'order_date': order_date.strftime('%d.%m.%Y'), 'enrollment_date': enrollment_date.strftime('%d.%m.%Y'), 'order_number': order_number, 'group_name': group.name, 'basis_documents': basis_documents})


def create_document_from_template(template_path, order_date, enrollment_date, order_number, group, applications):
    """Создание документа из шаблона для группы детей"""
    try:
        doc = DocxTemplate(template_path)
        
        children_data = []
        for i, app in enumerate(applications, 1):
            # Склоняем ФИО каждого ребенка
            declined_name = decline_fio_accusative(app.child_full_name)
            
            # Данные свидетельства
            birth_certificate = f"{app.birth_certificate_series or ''} {app.birth_certificate_number or ''}".strip()
            if not birth_certificate:
                birth_certificate = 'не указано'
            
            children_data.append({
                'number': i,
                'full_name': app.child_full_name,
                'full_name_declined': declined_name,
                'birth_date': app.child_birth_date.strftime('%d.%m.%Y'),
                'birth_certificate': birth_certificate,
                'birth_certificate_issued_by': app.birth_certificate_issued_by or 'не указано',
                'birth_certificate_issue_date': app.birth_certificate_issue_date.strftime('%d.%m.%Y') if app.birth_certificate_issue_date else ''
            })
        
        context = {
            'order_number': order_number,
            'order_date': order_date.strftime('%d.%m.%Y'),
            'enrollment_date': enrollment_date.strftime('%d.%m.%Y'),
            'group_name': group.name,
            'children': children_data
        }
        
        doc.render(context)
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        filename = f"Приказ_о_зачислении_{order_number}.docx"
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
        
    except Exception as e:
        raise e

@login_required
@director_required
def create_order_from_contract(request, contract_id):
    """Создание приказа о зачислении на основе подписанного договора"""
    contract = get_object_or_404(EducationContract, id=contract_id, status='signed')
    
    if hasattr(contract, 'enrollment_order') and contract.enrollment_order:
        messages.error(request, 'Для этого договора уже создан приказ')
        return redirect('orders:order_detail', pk=contract.enrollment_order.pk)
    
    if request.method == 'POST':
        group_id = request.POST.get('group_id')
        group = get_object_or_404(Group, id=group_id)
        enrollment_date = request.POST.get('enrollment_date', contract.enrollment_date)
        
        if isinstance(enrollment_date, str):
            try:
                enrollment_date = datetime.strptime(enrollment_date, '%Y-%m-%d').date()
            except:
                enrollment_date = contract.enrollment_date
        
        with transaction.atomic():
            child = Child.objects.create(
                full_name=contract.child_full_name,
                birth_date=contract.child_birth_date,
                gender=contract.application.child_gender if contract.application and hasattr(contract.application, 'child_gender') else 'M',
                registration_address=contract.application.registration_address if contract.application else '',
                actual_address=contract.application.actual_address if contract.application else '',
                group=group,
                enrollment_date=enrollment_date,
                is_active=True
            )
            
            ChildParent.objects.create(
                child=child,
                parent=contract.parent,
                relation='mother',
                is_primary=True
            )
            
            contract.child = child
            contract.save()
            
            order = EnrollmentOrder(
                contract=contract,
                child=child,
                group=group,
                enrollment_date=enrollment_date,
                basis_documents=f"На основании договора №{contract.contract_number} от {contract.registration_date.strftime('%d.%m.%Y')}",
                created_by=request.user,
                status='issued',
                order_date=date.today(),
                application=contract.application
            )
            order.save()
            
            contract.status = 'active'
            contract.save()
            
            if contract.application:
                contract.application.status = 'enrolled'
                contract.application.save()
            
            messages.success(request, f'Приказ №{order.order_number} создан. Ребенок {child.full_name} зачислен в группу {group.name}!')
            return redirect('orders:order_detail', pk=order.pk)
    
    groups = Group.objects.all()
    return render(request, 'orders/create_from_contract.html', {
        'contract': contract,
        'groups': groups,
        'today': date.today()
    })