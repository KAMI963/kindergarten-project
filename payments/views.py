from decimal import Decimal
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.http import JsonResponse, HttpResponse
from django.db.models import Q, Sum
from children.models import Child
from django.db import models as db_models
from .models import BankCard, GroupServiceAssignment, Payment, ServiceAttendance, Tariff, PaymentCalculation, AdditionalService, ServiceEnrollment, ChildAccount
from .forms import AdditionalServiceForm, ServiceEnrollmentForm
import json
import io
import os
from django.conf import settings
from docx import Document
from docx.shared import Pt
from datetime import date, datetime
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from celery import shared_task
from django.db import models
from children.models import Child

from django.contrib.auth.decorators import user_passes_test
from django.core.exceptions import PermissionDenied

# Декоратор для проверки роли директора
def director_required(function=None, redirect_field_name=None, login_url=None):
    """
    Декоратор, который проверяет, является ли пользователь директором
    """
    actual_decorator = user_passes_test(
        lambda u: u.is_authenticated and (u.is_superuser or u.is_staff or getattr(u, 'role', '') == 'director'),
        login_url=login_url,
        redirect_field_name=redirect_field_name
    )
    if function:
        return actual_decorator(function)
    return actual_decorator

def get_parent_children(user):
    """Получить детей родителя"""
    if not hasattr(user, 'parentprofile'):
        return Child.objects.none()
    
    try:
        return Child.objects.filter(
            parent_relations__parent=user.parentprofile,
            is_active=True
        ).distinct()
    except Exception as e:
        logger.error(f"Error getting parent children: {e}")
        return Child.objects.none()



import json
from collections import defaultdict


@login_required
def parent_finance_cabinet(request):
    """Личный кабинет «Финансы» (полная реализация по макету Doc1.docx)"""
    if request.user.role != 'parent':
        messages.error(request, 'Доступ только для родителей.')
        return redirect('dashboard')

    if not hasattr(request.user, 'parentprofile'):
        messages.warning(request, 'Пожалуйста, заполните анкету родителя')
        return redirect('accounts:parent_profile')

    children = get_parent_children(request.user)
    payments = Payment.objects.filter(child__in=children).select_related('child').order_by('-year', '-month')

    # === БАЛАНС И ОБЩАЯ СТАТИСТИКА ===
    total_charged = payments.aggregate(total=Sum('amount'))['total'] or Decimal('0')
    total_paid = payments.filter(status='completed').aggregate(total=Sum('amount'))['total'] or Decimal('0')
    current_balance = total_paid - total_charged

    today = date.today()
    current_month_payment = payments.filter(
        month=today.month, year=today.year, status__in=['pending', 'overdue']
    ).first()

    # === ДЕТАЛИЗАЦИЯ ПО МЕСЯЦАМ (реальные суммы) ===
    monthly_details = []
    month_names_full = [
        'Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
        'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь'
    ]

    for i in range(5, -1, -1):
        m = today.month - i
        y = today.year
        while m <= 0:
            m += 12
            y -= 1

        mp = payments.filter(month=m, year=y)
        charged = mp.aggregate(t=Sum('amount'))['t'] or Decimal('0')
        paid = mp.filter(status='completed').aggregate(t=Sum('amount'))['t'] or Decimal('0')

        if charged > 0 and paid >= charged:
            status = 'paid'
        elif charged > 0 and paid < charged:
            status = 'pending'
        else:
            status = 'empty'

        monthly_details.append({
            'name': month_names_full[m - 1],
            'year': y,
            'month_num': m,
            'charged': charged,
            'paid': paid,
            'status': status,
        })

    # === ДАННЫЕ ДЛЯ ГРАФИКОВ (безопасный JSON) ===
    chart_data = {
        'labels': [d['name'][:3] for d in monthly_details],
        'charged': [float(d['charged']) for d in monthly_details],
        'paid': [float(d['paid']) for d in monthly_details],
    }

    # === ИСТОРИЯ ПЛАТЕЖЕЙ С КАТЕГОРИЯМИ (Питание / Занятия) ===
    filter_child = request.GET.get('child_id')
    filter_month = request.GET.get('month')
    filter_year = request.GET.get('year')

    history_qs = payments.order_by('-year', '-month')
    if filter_child:
        history_qs = history_qs.filter(child_id=filter_child)
    if filter_month and filter_year:
        history_qs = history_qs.filter(month=int(filter_month), year=int(filter_year))

    history_list = []
    for p in history_qs[:30]:
        categories = []
        if p.base_amount > 0:
            categories.append({'name': 'Питание', 'amount': p.base_amount})
        if p.additional_services_amount > 0:
            categories.append({'name': 'Занятия', 'amount': p.additional_services_amount})
        if not categories:
            categories.append({'name': 'Прочее', 'amount': p.amount})

        history_list.append({
            'payment': p,
            'categories': categories,
        })

    # Группировка по детям для табов
    history_by_child = defaultdict(list)
    for item in history_list:
        history_by_child[item['payment'].child.full_name].append(item)

    saved_cards = BankCard.objects.filter(user=request.user, is_active=True)

    context = {
        'children': children,
        'total_charged': total_charged,
        'total_paid': total_paid,
        'current_balance': current_balance,
        'current_month_payment': current_month_payment,
        'monthly_details': monthly_details,
        'chart_data': json.dumps(chart_data),
        'history_by_child': dict(history_by_child),
        'all_history': history_list,
        'saved_cards': saved_cards,
        'filter_child': filter_child,
        'filter_month': filter_month,
        'filter_year': filter_year,
        'month_names_list': month_names_full,
    }
    return render(request, 'payments/parent_finance.html', context)
    
def calculate_payment_amount(payment):
    """Рассчитать сумму платежа"""
    try:
        return payment.calculate_amount()
    except Exception as e:
        logger.error(f"Error calculating payment: {e}")
        return Decimal('0')    
    

@shared_task
def auto_calculate_all_payments():
    """
    Автоматический расчет всех платежей 25-го числа каждого месяца
    Запускать через celery beat или вручную
    """
    today = date.today()
    
    # Если сегодня 25-е или меньше (для теста)
    if today.day <= 25:
        month = today.month
        year = today.year
    else:
        # Если после 25-го, рассчитываем следующий месяц
        if today.month == 12:
            month = 1
            year = today.year + 1
        else:
            month = today.month + 1
            year = today.year
    
    children = Child.objects.filter(is_active=True)
    created_count = 0
    updated_count = 0
    
    for child in children:
        payment, created = Payment.objects.get_or_create(
            child=child,
            month=month,
            year=year,
            defaults={'amount': 0}
        )
        
        # Рассчитываем сумму на основе посещаемости
        amount = payment.calculate_amount()
        payment.save()
        
        # Сохраняем историю расчета
        PaymentCalculation.objects.create(
            payment=payment,
            details={
                'base_amount': float(payment.base_amount),
                'attendance_days': payment.attendance_days,
                'total_days': payment.total_days,
                'food_days': payment.food_days,
                'discount': float(payment.discount_amount),
                'calculated_at': timezone.now().isoformat()
            }
        )
        
        if created:
            created_count += 1
        else:
            updated_count += 1
    
    return f"Расчет выполнен: создано {created_count}, обновлено {updated_count} платежей"

@login_required
def payment_history(request):
    """История платежей"""
    if request.user.role == 'parent':
        children = get_parent_children(request.user)
        payments = Payment.objects.filter(child__in=children).order_by('-year', '-month')
    elif request.user.role == 'director':
        payments = Payment.objects.all().order_by('-year', '-month')
    else:
        payments = Payment.objects.none()
    
    payment_groups = {}
    for payment in payments:
        key = f"{payment.year}-{payment.month:02d}"
        if key not in payment_groups:
            payment_groups[key] = []
        payment_groups[key].append(payment)
    
    return render(request, 'payments/payment_history.html', {
        'payment_groups': payment_groups,
        'payments': payments,
        'completed_count': payments.filter(status='completed').count(),
        'pending_count': payments.filter(status='pending').count(),
        'total_amount': payments.aggregate(total=Sum('amount'))['total'] or 0,
    })



@login_required
def payment_calculation(request, child_id, month=None, year=None):
    """
    Страница расчета платежа
    """
    child = get_object_or_404(Child, id=child_id)
    
    if request.user.role == 'parent':
        if not hasattr(request.user, 'parentprofile'):
            messages.error(request, 'Профиль родителя не найден.')
            return redirect('payments:payment_history')
        
        is_authorized = child.parent_relations.filter(parent__user=request.user).exists()
        if not is_authorized:
            messages.error(request, 'У вас нет доступа к этому ребенку.')
            return redirect('payments:payment_history')
    
    if not month or not year:
        now = timezone.now()
        month = now.month
        year = now.year
    
    payment, created = Payment.objects.get_or_create(
        child=child,
        month=month,
        year=year,
        defaults={'amount': 0}
    )
    
    if created or request.GET.get('recalculate'):
        amount = payment.calculate_amount()
        payment.save()
        
        calculation_details = {
            'base_amount': float(payment.base_amount),
            'attendance_days': payment.attendance_days,
            'total_days': payment.total_days,
            'food_days': payment.food_days,
            'discount': float(payment.discount_amount),  
            'additional_services_amount': float(payment.additional_services_amount),
            'calculated_at': timezone.now().isoformat()
        }
        
        PaymentCalculation.objects.create(
            payment=payment,
            details=calculation_details
        )
        
        if created:
            messages.success(request, f'Создан новый платеж на {amount} ₽')
        else:
            messages.info(request, f'Платеж пересчитан: {amount} ₽')
    
    calculations = PaymentCalculation.objects.filter(payment=payment).order_by('-calculation_date')
    
    return render(request, 'payments/payment_calculation.html', {
        'child': child,
        'payment': payment,
        'calculations': calculations,
        'month': month,
        'year': year
    })

from django.views.decorators.csrf import csrf_exempt
import json

@login_required
@csrf_exempt
def add_card_api(request):
    """API для добавления банковской карты"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Метод не разрешен'}, status=405)
    
    try:
        data = json.loads(request.body)
        
        card_holder = data.get('card_holder')
        card_number = data.get('card_number')
        card_type = data.get('card_type', 'visa')
        expiry_month = data.get('expiry_month')
        expiry_year = data.get('expiry_year')
        set_as_default = data.get('set_as_default', False)
        
        if not all([card_holder, card_number, expiry_month, expiry_year]):
            return JsonResponse({'success': False, 'error': 'Заполните все поля'}, status=400)
        
        # Проверка длины номера карты
        if len(card_number) != 16:
            return JsonResponse({'success': False, 'error': 'Номер карты должен содержать 16 цифр'}, status=400)
        
        # Маскируем номер карты
        masked_number = '•••• ' + card_number[-4:]
        
        # Если устанавливаем как основную, снимаем флаг с других карт
        if set_as_default:
            BankCard.objects.filter(user=request.user, is_active=True).update(is_default=False)
        
        card = BankCard.objects.create(
            user=request.user,
            card_type=card_type,
            card_number=card_number[-4:],
            card_number_masked=masked_number,
            card_holder=card_holder.upper(),
            expiry_month=expiry_month,
            expiry_year=expiry_year,
            is_default=set_as_default,
            is_active=True
        )
        
        return JsonResponse({
            'success': True, 
            'card_id': card.id,
            'masked_number': masked_number,
            'message': 'Карта успешно добавлена'
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@csrf_exempt
def delete_card_api(request, card_id):
    """API для удаления карты"""
    try:
        card = BankCard.objects.get(id=card_id, user=request.user)
        card.is_active = False
        card.save()
        return JsonResponse({'success': True, 'message': 'Карта удалена'})
    except BankCard.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Карта не найдена'}, status=404)


@login_required
@csrf_exempt
def set_default_card_api(request, card_id):
    """API для установки карты по умолчанию"""
    try:
        # Снимаем флаг со всех карт
        BankCard.objects.filter(user=request.user).update(is_default=False)
        # Устанавливаем выбранную карту
        card = BankCard.objects.get(id=card_id, user=request.user)
        card.is_default = True
        card.save()
        return JsonResponse({'success': True, 'message': 'Основная карта изменена'})
    except BankCard.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Карта не найдена'}, status=404)


@login_required
@csrf_exempt
def set_default_card_api(request, card_id):
    """API для установки карты по умолчанию"""
    try:
        BankCard.objects.filter(user=request.user).update(is_default=False)
        card = BankCard.objects.get(id=card_id, user=request.user)
        card.is_default = True
        card.save()
        return JsonResponse({'success': True, 'message': 'Основная карта изменена'})
    except BankCard.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Карта не найдена'}, status=404)


@login_required
def payment_page(request, child_id):
    """Страница оплаты с виртуальной картой"""
    child = get_object_or_404(Child, id=child_id)
    
    if request.user.role == 'parent':
        if not child.parent_relations.filter(parent__user=request.user).exists():
            messages.error(request, 'У вас нет доступа к этому ребенку.')
            return redirect('payments:parent_finance')
    
    now = timezone.now()
    payment, created = Payment.objects.get_or_create(
        child=child, month=now.month, year=now.year, defaults={'amount': 0}
    )
    
    if created or payment.amount == 0:
        payment.calculate_amount()
        payment.save()
    
    saved_cards = BankCard.objects.filter(user=request.user, is_active=True)
    
    recent_transactions = Payment.objects.filter(
        child__parent_relations__parent__user=request.user, status='completed'
    ).order_by('-paid_at')[:5]
    
    formatted_transactions = []
    for trans in recent_transactions:
        formatted_transactions.append({
            'description': f"Оплата за {trans.get_month_name()} {trans.year} - {trans.child.full_name}",
            'date': trans.paid_at or trans.created_at,
            'amount': -float(trans.amount),
            'type': 'expense', 'icon': 'credit-card', 'status': 'Оплачено'
        })
    
    total_balance = Payment.objects.filter(
        child__parent_relations__parent__user=request.user, status='completed'
    ).aggregate(total=models.Sum('amount'))['total'] or 0
    
    pending_amount = Payment.objects.filter(
        child__parent_relations__parent__user=request.user, status='pending'
    ).aggregate(total=models.Sum('amount'))['total'] or 0
    
    payment_progress = 100 if payment.status == 'completed' else 0
    refund_percentage = round(float(payment.discount_amount) / float(payment.base_amount) * 100, 1) if payment.base_amount > 0 else 0
    
    return render(request, 'payments/credit_cards_payment.html', {
        'child': child, 'payment': payment, 'saved_cards': saved_cards,
        'recent_transactions': formatted_transactions, 'total_balance': total_balance,
        'pending_amount': pending_amount, 'available_balance': total_balance - pending_amount,
        'payment_progress': payment_progress, 'refund_percentage': refund_percentage,
    })

@login_required
def simulate_payment(request, payment_id):
    """Имитация оплаты"""
    payment = get_object_or_404(Payment, id=payment_id)
    
    if request.user.role == 'parent':
        children_ids = get_parent_children(request.user).values_list('id', flat=True)
        if payment.child.id not in children_ids:
            messages.error(request, 'У вас нет доступа к этому платежу.')
            return redirect('payments:parent_finance')
    
    if request.method == 'POST':
        payment.status = 'completed'
        payment.paid_at = timezone.now()
        payment.transaction_id = f"TXN-{payment.id}-{int(timezone.now().timestamp())}"
        payment.payment_method = request.POST.get('payment_method', 'Банковская карта')
        payment.save()
        
        messages.success(request, f'Оплата на сумму {payment.amount} ₽ успешно проведена! Ожидайте возврат компенсации.')
        return redirect('payments:parent_finance')
    
    return render(request, 'payments/payment_checkout.html', {'payment': payment, 'child': payment.child})


@login_required
def recalculate_payment_ajax(request):
    """
    AJAX пересчет платежа
    """
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            payment_id = data.get('payment_id')
            
            payment = get_object_or_404(Payment, id=payment_id)
            
            if request.user.role == 'parent':
                if not payment.child.parent_relations.filter(parent__user=request.user).exists():
                    return JsonResponse({'success': False, 'error': 'Нет доступа'})
            
            amount = payment.calculate_amount()
            payment.save()
            
            calculation_details = {
                'base_amount': float(payment.base_amount),
                'attendance_days': payment.attendance_days,
                'total_days': payment.total_days,
                'food_days': payment.food_days,
                'discount': float(payment.discount_amount),
                'additional_services_amount': float(payment.additional_services_amount),
                'calculated_at': timezone.now().isoformat()
            }
            
            PaymentCalculation.objects.create(
                payment=payment,
                details=calculation_details
            )
            
            return JsonResponse({
                'success': True,
                'amount': float(payment.amount),
                'base_amount': float(payment.base_amount),
                'attendance_days': payment.attendance_days,
                'total_days': payment.total_days,
                'food_days': payment.food_days,
                'discount': float(payment.discount_amount),
                'additional_services_amount': float(payment.additional_services_amount)
            })
            
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})
    
    return JsonResponse({'success': False, 'error': 'Неверный запрос'})


@login_required
def payment_success(request):
    """
    Страница успешной оплаты
    """
    return render(request, 'payments/payment_success.html', {
        'now': timezone.now()
    })


@login_required
def credit_cards_payment_direct(request):
    """Страница оплаты с виртуальной картой - выбор ребенка"""
    if request.user.role != 'parent':
        messages.error(request, 'Доступ только для родителей')
        return redirect('dashboard')
    
    if not hasattr(request.user, 'parentprofile'):
        messages.warning(request, 'Пожалуйста, заполните анкету родителя')
        return redirect('accounts:parent_profile')
    
    children = get_parent_children(request.user)
    
    if not children.exists():
        messages.warning(request, 'У вас нет детей в системе')
        return redirect('dashboard')
    
    if children.count() == 1:
        return redirect('payments:payment_page', child_id=children.first().id)
    
    if request.method == 'POST':
        child_id = request.POST.get('child_id')
        if child_id:
            return redirect('payments:payment_page', child_id=child_id)
        else:
            messages.error(request, 'Выберите ребенка')
    
    now = timezone.now()
    children_with_payments = []
    
    for child in children:
        payment, created = Payment.objects.get_or_create(
            child=child,
            month=now.month,
            year=now.year,
            defaults={'amount': 0}
        )
        if payment.amount == 0:
            calculate_payment_amount(payment)
            payment.save()
        
        children_with_payments.append({
            'child': child,
            'payment': payment
        })
    
    return render(request, 'payments/select_child_for_payment.html', {
        'children_with_payments': children_with_payments
    })


@login_required
def credit_cards_payment_page(request, payment_id):
    """
    Страница оплаты банковской картой
    """
    payment = get_object_or_404(Payment, id=payment_id)
    
    # Проверка доступа
    if request.user.role == 'parent':
        if not payment.child.parent_relations.filter(parent__user=request.user).exists():
            messages.error(request, 'У вас нет доступа к этому платежу.')
            return redirect('payments:payment_history')
    
    # Получаем реквизиты для оплаты
    requisites = {} 
    
    # Получаем данные для отображения (используем правильные имена полей)
    context = {
        'payment': payment,
        'requisites': requisites,
        'child': payment.child,
        'amount': payment.amount,
        'month_name': payment.get_month_name(),
        'year': payment.year,
        'base_amount': payment.base_amount,
        'discount_amount': payment.discount_amount,  # Исправлено: discount_amount вместо discount
        'discount_rate': payment.discount_rate,
        'additional_services_amount': payment.additional_services_amount,
        'attendance_days': payment.attendance_days,
        'total_days': payment.total_days,
    }
    
    return render(request, 'payments/credit_cards_payment.html', context)


@login_required
@director_required
def recalculate_month_payments(request, year=None, month=None):
    """Пересчет платежей за конкретный месяц"""
    # Поддержка как позиционных аргументов из URL, так и GET/POST параметров
    if not year or not month:
        year = request.GET.get('year') or request.POST.get('year')
        month = request.GET.get('month') or request.POST.get('month')

    if not year or not month:
        messages.error(request, 'Не указан месяц или год.')
        return redirect('payments:director_payments_dashboard')

    try:
        year = int(year)
        month = int(month)
    except (ValueError, TypeError):
        messages.error(request, 'Некорректные значения месяца/года.')
        return redirect('payments:director_payments_dashboard')

    payments = Payment.objects.filter(month=month, year=year)
    updated = 0
    for payment in payments:
        payment.calculate_amount()
        payment.save()
        PaymentCalculation.objects.create(
            payment=payment,
            details={
                'base_amount': float(payment.base_amount),
                'attendance_days': payment.attendance_days,
                'total_days': payment.total_days,
                'food_days': payment.food_days,
                'discount_amount': float(payment.discount_amount),
                'additional_services_amount': float(payment.additional_services_amount),
                'amount': float(payment.amount),
                'calculated_at': timezone.now().isoformat(),
            }
        )
        updated += 1

    messages.success(request, f'Пересчитано {updated} платежей за {month}/{year}')
    return redirect('payments:director_payments_dashboard')


@login_required
def tariff_management(request):
    """
    Управление тарифами (только для заведующей)
    """
    if request.user.role != 'director':
        messages.error(request, 'Доступ только для заведующей.')
        return redirect('dashboard')
    
    tariffs = Tariff.objects.all().order_by('-year', 'age_category')
    
    if request.method == 'POST':
        # Проверяем, это редактирование или создание
        tariff_id = request.POST.get('tariff_id')
        
        age_category = request.POST.get('age_category')
        year = request.POST.get('year')
        
        # Используем правильные имена полей
        total_amount = request.POST.get('total_amount')  # Общая сумма за месяц
        content_amount = request.POST.get('content_amount')  # Содержание
        food_amount = request.POST.get('food_amount')  # Питание
        
        try:
            year = int(year)
            total_amount = float(total_amount) if total_amount else 0
            content_amount = float(content_amount) if content_amount else 0
            food_amount = float(food_amount) if food_amount else 0
            
            # Рассчитываем дневные ставки (условно: 20 рабочих дней в месяце)
            working_days = 20
            daily_rate = content_amount / working_days if working_days > 0 else 0
            food_daily_rate = food_amount / working_days if working_days > 0 else 0
            
        except (ValueError, TypeError):
            messages.error(request, 'Пожалуйста, введите корректные числовые значения.')
            return redirect('payments:tariff_management')
        
        if tariff_id:
            # Редактирование существующего тарифа
            tariff = get_object_or_404(Tariff, id=tariff_id)
            tariff.total_amount = total_amount
            tariff.content_amount = content_amount
            tariff.food_amount = food_amount
            tariff.daily_rate = daily_rate
            tariff.food_daily_rate = food_daily_rate
            tariff.save()
            messages.success(request, 'Тариф успешно обновлен!')
        else:
            # Создание нового тарифа
            tariff, created = Tariff.objects.update_or_create(
                age_category=age_category,
                year=year,
                defaults={
                    'total_amount': total_amount,
                    'content_amount': content_amount,
                    'food_amount': food_amount,
                    'daily_rate': daily_rate,
                    'food_daily_rate': food_daily_rate,
                }
            )
            if created:
                messages.success(request, 'Тариф успешно создан!')
            else:
                messages.success(request, 'Тариф успешно обновлен!')
        
        return redirect('payments:tariff_management')
    
    return render(request, 'payments/tariff_management.html', {
        'tariffs': tariffs
    })
    
    
@login_required
@director_required
def recalculate_all_payments(request):
    """Пересчет всех платежей (только для заведующей)"""
    payments = Payment.objects.all()
    updated_count = 0
    
    for payment in payments:
        old_amount = payment.amount
        payment.calculate_amount()
        payment.save()
        
        # Сохраняем историю расчета
        PaymentCalculation.objects.create(
            payment=payment,
            details={
                'base_amount': float(payment.base_amount),
                'attendance_days': payment.attendance_days,
                'total_days': payment.total_days,
                'food_days': payment.food_days,
                'discount': float(payment.discount_amount),
                'additional_services_amount': float(payment.additional_services_amount),
                'old_amount': float(old_amount),
                'new_amount': float(payment.amount),
                'calculated_at': timezone.now().isoformat()
            }
        )
        updated_count += 1
    
    messages.success(request, f'Пересчитано {updated_count} платежей!')
    return redirect('payments:payment_history')

@login_required
def generate_receipt(request, payment_id):
    """
    Генерация квитанции из шаблона DOCX с шрифтом 9
    """
    payment = get_object_or_404(Payment, id=payment_id)
    
    # Исправьте эту часть - проверка доступа для родителя
    if request.user.role == 'parent':
        # Используем правильное имя поля 'parent_relations'
        if not payment.child.parent_relations.filter(parent__user=request.user).exists():
            messages.error(request, 'У вас нет доступа к этому платежу.')
            return redirect('payments:payment_history')
    
    template_path = os.path.join(settings.BASE_DIR, 'templates', 'payments', 'Квитанция.docx')
    
    # Если файла шаблона нет, создаем простую квитанцию
    if not os.path.exists(template_path):
        # Создаем простой документ если шаблона нет
        doc = Document()
        doc.add_heading('КВИТАНЦИЯ', 0)
        doc.add_paragraph(f'Номер: R-{payment.id}-{payment.paid_at.strftime("%Y%m%d") if payment.paid_at else timezone.now().strftime("%Y%m%d")}')
        doc.add_paragraph(f'Плательщик: {request.user.get_full_name() if request.user.role == "parent" else "Родитель"}' if request.user.role == 'parent' else f'Плательщик: {payment.child.full_name}')
        doc.add_paragraph(f'Ребенок: {payment.child.full_name}')
        doc.add_paragraph(f'Период: {payment.get_month_name()} {payment.year}')
        doc.add_paragraph(f'Сумма: {payment.amount} ₽')
        
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="receipt_{payment.child.full_name}_{payment.month}_{payment.year}.docx"'
        return response
    
    doc = Document(template_path)
    
    payment_date = payment.paid_at if payment.paid_at else timezone.now()
    content_amount = payment.daily_rate * payment.attendance_days
    food_amount = payment.food_daily_rate * payment.food_days
    
    # Получаем имя плательщика
    if request.user.role == 'parent':
        payer_name = request.user.get_full_name() or request.user.username
    else:
        # Исправлено: используем правильное имя поля 'parent_relations'
        parents = payment.child.parent_relations.all()
        if parents.exists():
            payer_name = parents.first().parent.user.get_full_name() or parents.first().parent.user.username
        else:
            payer_name = "Не указан"
    
    # Функция для установки шрифта 9
    def set_font_size_9(element):
        if hasattr(element, 'runs'):
            for run in element.runs:
                run.font.size = Pt(9)
    
    # Обрабатываем параграфы
    for paragraph in doc.paragraphs:
        original_text = paragraph.text
        
        if '{{Номер_квитанции}}' in original_text:
            paragraph.text = original_text.replace('{{Номер_квитанции}}', f'R-{payment.id}-{payment_date.strftime("%Y%m%d")}')
        
        if '{{Номер}}' in original_text:
            paragraph.text = original_text.replace('{{Номер}}', str(payment.child.id))
        
        if '{{Плательщик}}}' in original_text:
            paragraph.text = original_text.replace('{{Плательщик}}}', payer_name)
        
        if '{{Ребенок}}}' in original_text:
            paragraph.text = original_text.replace('{{Ребенок}}}', payment.child.full_name)
        
        if '{{Плата}}' in original_text:
            paragraph.text = original_text.replace('{{Плата}}', f'{content_amount:.2f}')
        
        if '{{Питание}}' in original_text:
            paragraph.text = original_text.replace('{{Питание}}', f'{food_amount:.2f}')
        
        if '{{Сумма}}' in original_text:
            paragraph.text = original_text.replace('{{Сумма}}', f'{payment.amount:.2f}')
        
        if '{{Дата_оплаты}}' in original_text:
            paragraph.text = original_text.replace('{{Дата_оплаты}}', payment_date.strftime('%d.%m.%Y'))
        
        if '{{Лицевой_счет}}' in original_text:
            account_number = getattr(payment.child, 'account', None)
            account_str = account_number.account_number if account_number else "Не указан"
            paragraph.text = original_text.replace('{{Лицевой_счет}}', account_str)
        
        set_font_size_9(paragraph)
    
    # Обрабатываем таблицы
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                original_text = cell.text
                
                if '{{Номер_квитанции}}' in original_text:
                    cell.text = original_text.replace('{{Номер_квитанции}}', f'R-{payment.id}-{payment_date.strftime("%Y%m%d")}')
                
                if '{{Номер}}' in original_text:
                    cell.text = original_text.replace('{{Номер}}', str(payment.child.id))
                
                if '{{Плательщик}}}' in original_text:
                    cell.text = original_text.replace('{{Плательщик}}}', payer_name)
                
                if '{{Ребенок}}}' in original_text:
                    cell.text = original_text.replace('{{Ребенок}}}', payment.child.full_name)
                
                if '{{Плата}}' in original_text:
                    cell.text = original_text.replace('{{Плата}}', f'{content_amount:.2f}')
                
                if '{{Питание}}' in original_text:
                    cell.text = original_text.replace('{{Питание}}', f'{food_amount:.2f}')
                
                if '{{Сумма}}' in original_text:
                    cell.text = original_text.replace('{{Сумма}}', f'{payment.amount:.2f}')
                
                if '{{Дата_оплаты}}' in original_text:
                    cell.text = original_text.replace('{{Дата_оплаты}}', payment_date.strftime('%d.%m.%Y'))
                
                if '{{Лицевой_счет}}' in original_text:
                    account_number = getattr(payment.child, 'account', None)
                    account_str = account_number.account_number if account_number else "Не указан"
                    cell.text = original_text.replace('{{Лицевой_счет}}', account_str)
                
                # Устанавливаем шрифт для всех параграфов в ячейке
                for paragraph in cell.paragraphs:
                    set_font_size_9(paragraph)
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    response = HttpResponse(
        file_stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="receipt_{payment.child.full_name}_{payment.month}_{payment.year}.docx"'
    
    return response


@login_required
def services_catalog(request):
    """
    Каталог дополнительных услуг (для родителей)
    """
    services = AdditionalService.objects.filter(is_active=True)
    
    # Получаем детей родителя - используем правильное имя поля 'parent_relations'
    children_list = []
    if request.user.role == 'parent' and hasattr(request.user, 'parentprofile'):
        children_list = Child.objects.filter(parent_relations__parent=request.user.parentprofile)
    
    # Фильтрация по возрасту ребенка
    child_id = request.GET.get('child_id')
    child = None
    available_services = services
    
    if child_id and request.user.role == 'parent':
        try:
            child = Child.objects.get(id=child_id)
            # Проверяем, что ребенок принадлежит родителю
            if child in children_list:
                child_age = child.get_age() if hasattr(child, 'get_age') else 3
                available_services = services.filter(
                    age_min__lte=child_age,
                    age_max__gte=child_age
                )
        except Child.DoesNotExist:
            pass
    
    # Получаем уже записанные услуги
    enrolled_services = {}
    if request.user.role == 'parent' and child:
        enrollments = ServiceEnrollment.objects.filter(
            child=child,
            status__in=['pending', 'approved']
        ).select_related('service')
        for enrollment in enrollments:
            enrolled_services[enrollment.service.id] = enrollment.status
    
    return render(request, 'payments/services_catalog.html', {
        'services': available_services,
        'child': child,
        'enrolled_services': enrolled_services,
        'children': children_list
    })


@login_required
def enroll_in_service(request, service_id):
    """
    Подача заявки на дополнительную услугу
    """
    service = get_object_or_404(AdditionalService, id=service_id, is_active=True)
    
    # Получаем детей родителя
    children = []
    if request.user.role == 'parent' and hasattr(request.user, 'parentprofile'):
        children = Child.objects.filter(parent_relations__parent=request.user.parentprofile)
    
    if request.method == 'POST':
        # Получаем child_id из POST данных
        child_id = request.POST.get('child_id')
        start_date = request.POST.get('start_date')
        notes = request.POST.get('notes', '')
        
        # Проверяем, что child_id выбран
        if not child_id:
            messages.error(request, 'Пожалуйста, выберите ребенка.')
            return redirect('payments:enroll_in_service', service_id=service_id)
        
        try:
            child = Child.objects.get(id=child_id)
        except Child.DoesNotExist:
            messages.error(request, 'Ребенок не найден.')
            return redirect('payments:enroll_in_service', service_id=service_id)
        
        # Проверка доступа - ребенок должен принадлежать родителю
        if request.user.role == 'parent':
            if child not in children:
                messages.error(request, 'У вас нет доступа к этому ребенку.')
                return redirect('payments:enroll_in_service', service_id=service_id)
        
        # Проверка возраста ребенка
        child_age = child.get_age() if hasattr(child, 'get_age') else 3
        if not (service.age_min <= child_age <= service.age_max):
            messages.error(request, f'Ребенок не подходит по возрасту для этой услуги. Возраст должен быть от {service.age_min} до {service.age_max} лет.')
            return redirect('payments:enroll_in_service', service_id=service_id)
        
        # Проверка на уже существующую заявку
        existing = ServiceEnrollment.objects.filter(
            child=child,
            service=service,
            status__in=['pending', 'approved']
        ).first()
        
        if existing:
            messages.warning(request, f'Заявка на услугу "{service.name}" уже подана.')
            return redirect('payments:services_catalog')
        
        # Создаем заявку
        from datetime import date
        enrollment = ServiceEnrollment(
            child=child,
            service=service,
            start_date=start_date if start_date else date.today(),
            notes=notes,
            status='pending'
        )
        enrollment.save()
        
        messages.success(request, f'Заявка на услугу "{service.name}" успешно подана! Статус: На рассмотрении')
        return redirect('payments:my_enrollments')
    
    # GET запрос - показываем форму
    # Проверка возраста ребенка для каждого ребенка
    children_with_age = []
    for child in children:
        child_age = child.get_age() if hasattr(child, 'get_age') else 3
        is_age_ok = service.age_min <= child_age <= service.age_max
        children_with_age.append({
            'child': child,
            'age': child_age,
            'is_age_ok': is_age_ok
        })
    
    return render(request, 'payments/enroll_service.html', {
        'service': service,
        'children': children_with_age
    })

@login_required
def my_enrollments(request):
    """
    Мои заявки на дополнительные услуги (для родителей)
    """
    if request.user.role != 'parent' or not hasattr(request.user, 'parentprofile'):
        messages.error(request, 'Доступ только для родителей.')
        return redirect('dashboard')
    
    # Получаем детей родителя - используем правильное имя поля
    children = Child.objects.filter(parent_relations__parent=request.user.parentprofile)
    
    enrollments = ServiceEnrollment.objects.filter(
        child__in=children
    ).select_related('child', 'service').order_by('-enrollment_date')
    
    return render(request, 'payments/my_enrollments.html', {
        'enrollments': enrollments,
        'children': children
    })


@login_required
def group_service_enrollments(request):
    """
    Просмотр заявок на услуги для воспитателя/заведующей
    """
    if request.user.role not in ['director', 'teacher']:
        messages.error(request, 'Доступ запрещен.')
        return redirect('dashboard')
    
    enrollments = ServiceEnrollment.objects.select_related('child', 'service', 'child__group').order_by('-enrollment_date')
    
    if request.user.role == 'teacher' and hasattr(request.user, 'employee_profile'):
        # Воспитатель видит только свою группу
        group = getattr(request.user.employee_profile, 'group', None)
        if group:
            enrollments = enrollments.filter(child__group=group)
    
    status_filter = request.GET.get('status')
    if status_filter:
        enrollments = enrollments.filter(status=status_filter)
    
    service_filter = request.GET.get('service')
    if service_filter:
        enrollments = enrollments.filter(service_id=service_filter)
    
    return render(request, 'payments/group_enrollments.html', {
        'enrollments': enrollments,
        'status_choices': ServiceEnrollment.STATUS_CHOICES,
        'services': AdditionalService.objects.filter(is_active=True),
        'current_status': status_filter,
        'current_service': service_filter
    })


@login_required
def update_enrollment_status(request, enrollment_id):
    """
    Изменение статуса заявки (для заведующей/воспитателя)
    """
    if request.user.role not in ['director', 'teacher']:
        messages.error(request, 'Доступ запрещен.')
        return redirect('dashboard')
    
    enrollment = get_object_or_404(ServiceEnrollment, id=enrollment_id)
    
    if request.method == 'POST':
        new_status = request.POST.get('status')
        notes = request.POST.get('notes', '')
        
        if new_status in dict(ServiceEnrollment.STATUS_CHOICES):
            enrollment.status = new_status
            if notes:
                enrollment.notes = notes
            enrollment.save()
            
            messages.success(request, f'Статус заявки изменен на "{enrollment.get_status_display()}"')
    
    return redirect('payments:group_enrollments')


@login_required
def director_payments_dashboard(request):
    """
    Дашборд оплат для заведующей - мониторинг оплат и задолженностей
    """
    if request.user.role != 'director':
        messages.error(request, 'Доступ только для заведующей.')
        return redirect('dashboard')
    
    now = timezone.now()
    current_month = now.month
    current_year = now.year
    today = date.today()
    
    # Все платежи за текущий месяц
    current_payments = Payment.objects.filter(month=current_month, year=current_year).select_related('child')
    
    # Статистика
    total_payments = current_payments.count()
    paid_count = current_payments.filter(status='completed').count()
    pending_count = current_payments.filter(status='pending').count()
    
    # Подсчет просроченных (статус pending и дата больше 15 числа месяца)
    overdue_count = 0
    for payment in current_payments:
        if payment.status == 'pending' and payment.is_overdue:
            overdue_count += 1
    
    total_amount_due = current_payments.aggregate(total=Sum('amount'))['total'] or 0
    total_amount_paid = current_payments.filter(status='completed').aggregate(total=Sum('amount'))['total'] or 0
    
    # Прогресс оплаты
    paid_percentage = (paid_count / total_payments * 100) if total_payments > 0 else 0
    
    # Группировка по группам
    payments_by_group = {}
    for payment in current_payments:
        group_name = payment.child.group.name if payment.child.group else 'Без группы'
        if group_name not in payments_by_group:
            payments_by_group[group_name] = {
                'total': 0,
                'paid': 0,
                'pending': 0,
                'amount_due': 0,
                'amount_paid': 0
            }
        payments_by_group[group_name]['total'] += 1
        if payment.status == 'completed':
            payments_by_group[group_name]['paid'] += 1
            payments_by_group[group_name]['amount_paid'] += float(payment.amount)
        else:
            payments_by_group[group_name]['pending'] += 1
            payments_by_group[group_name]['amount_due'] += float(payment.amount)
    
    # Список должников (просрочка после 15 числа)
    debtors = []
    for payment in current_payments:
        if payment.status != 'completed' and payment.is_overdue:
            days_overdue = (today - date(payment.year, payment.month, 15)).days
            debtors.append({
                'child': payment.child,
                'payment': payment,
                'amount': payment.amount,
                'days_overdue': days_overdue
            })
    
    return render(request, 'payments/director_payments.html', {
        'current_month': current_month,
        'current_year': current_year,
        'total_payments': total_payments,
        'paid_count': paid_count,
        'pending_count': pending_count,
        'overdue_count': overdue_count,
        'total_amount_due': total_amount_due,
        'total_amount_paid': total_amount_paid,
        'paid_percentage': round(paid_percentage, 1),
        'payments_by_group': payments_by_group,
        'debtors': debtors,
        'payments': current_payments
    })


@login_required
def parent_payments(request):
    """Личный кабинет оплат для родителя"""
    if request.user.role != 'parent':
        messages.error(request, 'Доступ только для родителей.')
        return redirect('dashboard')
    
    if not hasattr(request.user, 'parentprofile'):
        messages.warning(request, 'Пожалуйста, заполните анкету родителя')
        return redirect('accounts:parent_profile')
    
    children = get_parent_children(request.user)
    payments = Payment.objects.filter(child__in=children).select_related('child').order_by('-year', '-month')
    
    payments_by_child = {}
    for child in children:
        payments_by_child[child.id] = {
            'child': child,
            'payments': [],
            'total_paid': 0,
            'total_due': 0,
            'has_overdue': False
        }
    
    for payment in payments:
        child_data = payments_by_child.get(payment.child.id)
        if child_data:
            child_data['payments'].append(payment)
            if payment.status == 'completed':
                child_data['total_paid'] += float(payment.amount)
            else:
                child_data['total_due'] += float(payment.amount)
                if hasattr(payment, 'is_overdue') and payment.is_overdue:
                    child_data['has_overdue'] = True
    
    return render(request, 'payments/parent_payments.html', {
        'payments_by_child': payments_by_child.values(),
        'payments': payments
    })

    
@login_required
def manage_services(request):
    """
    Управление дополнительными услугами (только для заведующей)
    """
    if not (request.user.is_superuser or request.user.is_staff or getattr(request.user, 'role', '') == 'director'):
        messages.error(request, 'Доступ только для заведующей.')
        return redirect('dashboard')
    
    services = AdditionalService.objects.all().order_by('-is_active', 'name')
    
    # Получаем список преподавателей (пользователей с ролью teacher)
    from django.contrib.auth import get_user_model
    User = get_user_model()
    teachers = User.objects.filter(role='teacher').order_by('first_name', 'last_name')
    
    return render(request, 'payments/manage_services.html', {
        'services': services,
        'teachers': teachers,
        'title': 'Управление дополнительными услугами'
    })


@login_required
def service_create(request):
    """Создание кружка — отдельная страница"""
    if request.user.role != 'director' and not request.user.is_superuser:
        messages.error(request, 'Доступ только для заведующей.')
        return redirect('dashboard')

    if request.method == 'POST':
        form = AdditionalServiceForm(request.POST)
        if form.is_valid():
            service = form.save()
            messages.success(request, f'Кружок «{service.name}» успешно создан!')
            return redirect('payments:manage_services')
        else:
            # Показываем ошибки формы
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f'{field}: {error}')
    else:
        form = AdditionalServiceForm()

    return render(request, 'payments/service_form.html', {
        'form': form,
        'title': 'Создание кружка',
        'service': None,
    })


@login_required
def service_edit(request, pk):
    """Редактирование кружка — отдельная страница"""
    if request.user.role != 'director' and not request.user.is_superuser:
        messages.error(request, 'Доступ только для заведующей.')
        return redirect('dashboard')

    service = get_object_or_404(AdditionalService, pk=pk)

    if request.method == 'POST':
        form = AdditionalServiceForm(request.POST, instance=service)
        if form.is_valid():
            form.save()
            messages.success(request, f'Кружок «{service.name}» успешно обновлён!')
            return redirect('payments:manage_services')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f'{field}: {error}')
    else:
        form = AdditionalServiceForm(instance=service)

    return render(request, 'payments/service_form.html', {
        'form': form,
        'title': f'Редактирование: {service.name}',
        'service': service,
    })


@login_required
def service_delete(request, pk):
    """
    Удаление услуги
    """
    if not (request.user.is_superuser or request.user.is_staff or getattr(request.user, 'role', '') == 'director'):
        messages.error(request, 'Доступ только для заведующей.')
        return redirect('dashboard')
    
    service = get_object_or_404(AdditionalService, pk=pk)
    
    if request.method == 'POST':
        service_name = service.name
        service.delete()
        messages.success(request, f'Услуга "{service_name}" успешно удалена!')
        return redirect('payments:manage_services')
    
    return render(request, 'payments/service_confirm_delete.html', {
        'service': service
    })


@login_required
def service_toggle_active(request, pk):
    if not (request.user.is_superuser or request.user.is_staff or getattr(request.user, 'role', '') == 'director'):
        return JsonResponse({'error': 'Доступ запрещен'}, status=403)
    
    service = get_object_or_404(AdditionalService, pk=pk)
    service.is_active = not service.is_active
    service.save()
    
    # ✅ Возвращаем success для JS
    return JsonResponse({
        'success': True,
        'message': f'Услуга "{service.name}" {"активирована" if service.is_active else "деактивирована"}'
    })



@login_required
def teacher_services_dashboard(request):
    """
    Дашборд воспитателя - просмотр своих кружков и расписания
    """
    if request.user.role != 'teacher':
        messages.error(request, 'Доступ только для воспитателей.')
        return redirect('dashboard')
    
    # Получаем кружки, где пользователь является преподавателем
    my_services = AdditionalService.objects.filter(
        teacher_user=request.user,
        is_active=True
    )
    
    # Получаем назначения кружков на группу воспитателя
    # Проверяем, есть ли у сотрудника группа
    if hasattr(request.user, 'employee_profile'):
        employee = request.user.employee_profile
        # Пробуем разные варианты получения группы
        group = None
        
        # Вариант 1: если есть поле group
        if hasattr(employee, 'group'):
            group = employee.group
        # Вариант 2: если есть поле assigned_group
        elif hasattr(employee, 'assigned_group'):
            group = employee.assigned_group
        # Вариант 3: если есть связь через related_name
        elif hasattr(employee, 'groups') and employee.groups.exists():
            group = employee.groups.first()
        
        if group:
            # Проверяем, есть ли модель GroupServiceAssignment
            try:
                from payments.models import GroupServiceAssignment
                group_assignments = GroupServiceAssignment.objects.filter(
                    group=group,
                    is_active=True
                ).select_related('service')
                for assignment in group_assignments:
                    if assignment.service not in my_services:
                        my_services = my_services | AdditionalService.objects.filter(id=assignment.service.id)
            except:
                # Если модель не существует, пропускаем
                pass
    
    # Если нет кружков, показываем сообщение
    if not my_services.exists():
        return render(request, 'payments/teacher_dashboard.html', {
            'my_services': my_services,
            'enrollments_by_service': {},
            'today': date.today(),
            'no_services': True
        })
    
    # Заявки на кружки воспитателя
    enrollments = ServiceEnrollment.objects.filter(
        service__in=my_services,
        status='approved'
    ).select_related('child', 'service', 'child__group').order_by('service__name', 'child__full_name')
    
    # Группировка по кружкам
    enrollments_by_service = {}
    for enrollment in enrollments:
        if enrollment.service.id not in enrollments_by_service:
            enrollments_by_service[enrollment.service.id] = {
                'service': enrollment.service,
                'enrollments': []
            }
        enrollments_by_service[enrollment.service.id]['enrollments'].append(enrollment)
    
    return render(request, 'payments/teacher_dashboard.html', {
        'my_services': my_services,
        'enrollments_by_service': enrollments_by_service,
        'today': date.today()
    })


@login_required
def mark_service_attendance(request, service_id):
    """
    Страница отметки посещаемости по конкретному кружку.
    Показывает всех записанных детей и форму быстрой отметки за сегодня.
    """
    service = get_object_or_404(AdditionalService, id=service_id, is_active=True)
    
    # Проверка прав: воспитатель должен вести этот кружок
    if request.user.role != 'teacher':
        messages.error(request, 'Доступ только для воспитателей.')
        return redirect('dashboard')
    
    if service.teacher_user != request.user:
        messages.error(request, 'У вас нет доступа к этому кружку.')
        return redirect('payments:teacher_dashboard')
    
    # Все одобренные записи на этот кружок
    enrollments = ServiceEnrollment.objects.filter(
        service=service,
        status='approved'
    ).select_related('child', 'child__group').order_by('child__full_name')
    
    today = date.today()
    
    # Загружаем уже существующие отметки за сегодня
    existing_attendance = {}
    attendances = ServiceAttendance.objects.filter(
        enrollment__in=enrollments,
        date=today
    ).select_related('enrollment')
    for att in attendances:
        existing_attendance[att.enrollment_id] = att.status
    
    # Обработка POST — сохранение отметок
    if request.method == 'POST':
        saved_count = 0
        for enrollment in enrollments:
            status_key = f'status_{enrollment.id}'
            notes_key = f'notes_{enrollment.id}'
            status = request.POST.get(status_key)
            notes = request.POST.get(notes_key, '')
            
            if status and status in dict(ServiceAttendance.ABSENCE_REASONS):
                attendance, created = ServiceAttendance.objects.update_or_create(
                    enrollment=enrollment,
                    date=today,
                    defaults={
                        'status': status,
                        'notes': notes,
                        'marked_by': request.user,
                    }
                )
                saved_count += 1
        
        messages.success(request, f'Посещаемость за {today.strftime("%d.%m.%Y")} сохранена ({saved_count} чел.)')
        
        # Пересчитываем платежи за этот месяц для всех отмеченных детей
        for enrollment in enrollments:
            payment = Payment.objects.filter(
                child=enrollment.child,
                month=today.month,
                year=today.year
            ).first()
            if payment:
                payment.calculate_amount()
                payment.save()
        
        return redirect('payments:mark_service_attendance', service_id=service_id)
    
    return render(request, 'payments/mark_attendance.html', {
        'service': service,
        'enrollments': enrollments,
        'today': today,
        'existing_attendance': existing_attendance,
        'absence_reasons': ServiceAttendance.ABSENCE_REASONS,
    })


@login_required
def service_attendance_sheet(request, service_id):
    """Табель посещаемости кружка (для воспитателя)"""
    service = get_object_or_404(AdditionalService, id=service_id)
    
    # Проверка прав
    if request.user.role != 'teacher':
        messages.error(request, 'Доступ только для воспитателей.')
        return redirect('dashboard')
    
    if service.teacher_user != request.user:
        if hasattr(request.user, 'employee_profile') and request.user.employee_profile.group:
            is_assigned = GroupServiceAssignment.objects.filter(
                group=request.user.employee_profile.group,
                service=service,
                is_active=True
            ).exists()
            if not is_assigned:
                messages.error(request, 'У вас нет доступа к этому кружку.')
                return redirect('payments:teacher_dashboard')
    
    enrollments = ServiceEnrollment.objects.filter(
        service=service, status='approved'
    ).select_related('child')
    
    month = int(request.GET.get('month', date.today().month))
    year = int(request.GET.get('year', date.today().year))
    
    import calendar
    days_in_month = calendar.monthrange(year, month)[1]
    dates = [date(year, month, day) for day in range(1, days_in_month + 1)]
    
    attendance_data = {}
    for enrollment in enrollments:
        attendance_data[enrollment.id] = {}
        attendances = ServiceAttendance.objects.filter(
            enrollment=enrollment, date__month=month, date__year=year
        )
        for att in attendances:
            attendance_data[enrollment.id][att.date.day] = att
    
    # ✅ ДОБАВЛЕНО: список названий месяцев для шаблона
    month_names = [
        '', 'Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
        'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь'
    ]
    
    return render(request, 'payments/attendance_sheet.html', {
        'service': service,
        'enrollments': enrollments,
        'dates': dates,
        'month': month,
        'year': year,
        'attendance_data': attendance_data,
        'absence_reasons': ServiceAttendance.ABSENCE_REASONS,
        'month_names': month_names,  # ← новая переменная
    })


@login_required
def generate_service_contract(request, enrollment_id):
    """
    Генерация договора на дополнительную услугу
    """
    enrollment = get_object_or_404(ServiceEnrollment, id=enrollment_id)
    
    # Проверка доступа
    if request.user.role == 'parent':
        if enrollment.child.parent_relations.filter(parent__user=request.user).exists():
            # Родитель может скачать договор для своего ребенка
            pass
        else:
            messages.error(request, 'У вас нет доступа')
            return redirect('dashboard')
    elif request.user.role == 'director':
        pass  # Заведующая может скачать любой договор
    else:
        messages.error(request, 'Доступ запрещен')
        return redirect('dashboard')
    
    # Путь к шаблону
    template_path = os.path.join(settings.BASE_DIR, 'templates', 'payments', 'Договор_о_платных_услугах.docx')
    
    if not os.path.exists(template_path):
        messages.error(request, 'Шаблон договора не найден')
        return redirect('payments:my_enrollments')
    
    doc = Document(template_path)
    
    # Получаем данные для подстановки
    parent = enrollment.child.parent_relations.first().parent if enrollment.child.parent_relations.exists() else None
    parent_name = parent.user.get_full_name() if parent else "Родитель"
    
    replacements = {
        '{{Ребенок_ФИО}}': enrollment.child.full_name,
        '{{Ребенок_Дата_рождения}}': enrollment.child.birth_date.strftime('%d.%m.%Y') if enrollment.child.birth_date else '',
        '{{Родитель_ФИО}}': parent_name,
        '{{Услуга_название}}': enrollment.service.name,
        '{{Услуга_стоимость}}': str(enrollment.service.price_per_month),
        '{{Услуга_расписание}}': enrollment.service.schedule,
        '{{Услуга_преподаватель}}': enrollment.service.teacher,
        '{{Дата_начала}}': enrollment.start_date.strftime('%d.%m.%Y'),
        '{{Дата_заключения}}': date.today().strftime('%d.%m.%Y'),
    }
    
    for paragraph in doc.paragraphs:
        for key, value in replacements.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for key, value in replacements.items():
                    if key in cell.text:
                        cell.text = cell.text.replace(key, value)
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    response = HttpResponse(
        file_stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="dogovor_{enrollment.child.full_name}_{enrollment.service.name}.docx"'
    
    return response


@login_required
def sign_service_contract(request, enrollment_id):
    """
    Электронная подпись договора (кнопка "Ознакомлен и согласен")
    """
    if request.method != 'POST':
        messages.error(request, 'Метод не поддерживается')
        return redirect('payments:my_enrollments')
    
    enrollment = get_object_or_404(ServiceEnrollment, id=enrollment_id)
    
    # Проверка доступа - только родитель может подписать
    if request.user.role != 'parent':
        messages.error(request, 'Только родитель может подписать договор')
        return redirect('dashboard')
    
    if not enrollment.child.parent_relations.filter(parent__user=request.user).exists():
        messages.error(request, 'У вас нет доступа к этому ребенку')
        return redirect('dashboard')
    
    # Подписываем договор
    enrollment.contract_signed = True
    enrollment.contract_signed_at = timezone.now()
    enrollment.save()
    
    # Отправляем уведомление
    messages.success(request, 
        'Благодарим Вас за оставленную заявку! Вам необходимо подойти в детский сад для подписания необходимых документов.')
    
    return redirect('payments:my_enrollments')

import logging
import traceback
logger = logging.getLogger(__name__)

@login_required
def api_service_detail(request, pk):
    """API для получения данных кружка"""
    try:
        if request.user.role != 'director' and not request.user.is_superuser:
            return JsonResponse({'error': 'Нет доступа'}, status=403)

        service = get_object_or_404(AdditionalService, pk=pk)

        data = {
            'id': service.id,
            'name': service.name,
            'service_type': service.service_type,
            'service_type_display': service.get_service_type_display(),
            'description': service.description or '',
            'price_per_month': float(service.price_per_month),
            'price_per_lesson': float(service.price_per_lesson),
            'age_min': service.age_min,
            'age_max': service.age_max,
            'teacher': service.teacher or '',
            'teacher_user_id': service.teacher_user_id,
            'teacher_name': service.teacher_user.get_full_name() if service.teacher_user else '',
            'schedule': service.schedule or '',
            'program': service.program or '',
            'is_active': service.is_active,
            'max_students': service.max_students,
        }
        return JsonResponse(data)

    except Exception as e:
        logger.error(f"ERROR in api_service_detail: {e}")
        return JsonResponse({'error': str(e)}, status=500)


@login_required
def service_create_ajax(request):
    """Создание кружка через AJAX"""
    if request.user.role != 'director' and not request.user.is_superuser:
        return JsonResponse({'error': 'Нет доступа'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'error': 'Метод не поддерживается'}, status=405)

    try:
        teacher_user_id = request.POST.get('teacher_user_id') or None
        
        # Валидация обязательных полей
        name = request.POST.get('name', '').strip()
        if not name:
            return JsonResponse({'error': 'Название кружка обязательно'}, status=400)
        
        price_per_lesson = request.POST.get('price_per_lesson', '0')
        price_per_month = request.POST.get('price_per_month', '0')
        
        service = AdditionalService.objects.create(
            name=name,
            service_type=request.POST.get('service_type', 'circle'),
            description=request.POST.get('description', ''),
            price_per_month=Decimal(price_per_month) if price_per_month else Decimal('0'),
            price_per_lesson=Decimal(price_per_lesson) if price_per_lesson else Decimal('0'),
            age_min=int(request.POST.get('age_min', 3)),
            age_max=int(request.POST.get('age_max', 7)),
            teacher=request.POST.get('teacher', ''),
            teacher_user_id=int(teacher_user_id) if teacher_user_id else None,
            schedule=request.POST.get('schedule', ''),
            program=request.POST.get('program', ''),
            is_active=request.POST.get('is_active', 'true').lower() == 'true',
            max_students=int(request.POST.get('max_students', 10)),
        )
        return JsonResponse({'success': True, 'id': service.id, 'name': service.name})
    
    except ValueError as e:
        return JsonResponse({'error': f'Некорректные данные: {str(e)}'}, status=400)
    except Exception as e:
        logger.error(f"Error creating service: {e}")
        return JsonResponse({'error': str(e)}, status=400)


@login_required
def service_edit_ajax(request, pk):
    """Редактирование кружка через AJAX"""
    if request.user.role != 'director' and not request.user.is_superuser:
        return JsonResponse({'error': 'Нет доступа'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'error': 'Метод не поддерживается'}, status=405)

    try:
        service = get_object_or_404(AdditionalService, pk=pk)
        
        teacher_user_id = request.POST.get('teacher_user_id') or None
        
        price_per_lesson = request.POST.get('price_per_lesson', '0')
        price_per_month = request.POST.get('price_per_month', '0')
        
        service.name = request.POST.get('name', service.name).strip()
        service.service_type = request.POST.get('service_type', service.service_type)
        service.description = request.POST.get('description', '')
        service.price_per_month = Decimal(price_per_month) if price_per_month else Decimal('0')
        service.price_per_lesson = Decimal(price_per_lesson) if price_per_lesson else Decimal('0')
        service.age_min = int(request.POST.get('age_min', 3))
        service.age_max = int(request.POST.get('age_max', 7))
        service.teacher = request.POST.get('teacher', '')
        service.teacher_user_id = int(teacher_user_id) if teacher_user_id else None
        service.schedule = request.POST.get('schedule', '')
        service.program = request.POST.get('program', '')
        service.is_active = request.POST.get('is_active', 'true').lower() == 'true'
        service.max_students = int(request.POST.get('max_students', 10))
        
        # Валидация возрастов
        if service.age_min > service.age_max:
            return JsonResponse({'error': 'Минимальный возраст не может быть больше максимального'}, status=400)
        
        service.save()
        return JsonResponse({'success': True, 'id': service.id, 'name': service.name})
    
    except ValueError as e:
        return JsonResponse({'error': f'Некорректные данные: {str(e)}'}, status=400)
    except Exception as e:
        logger.error(f"Error editing service: {e}")
        return JsonResponse({'error': str(e)}, status=400)


@login_required
@require_http_methods(["POST"])
def service_delete_ajax(request, pk):
    """Удаление кружка через AJAX"""
    if request.user.role != 'director' and not request.user.is_superuser:
        return JsonResponse({'error': 'Нет доступа'}, status=403)
    
    service = get_object_or_404(AdditionalService, pk=pk)
    service.delete()
    return JsonResponse({'success': True})